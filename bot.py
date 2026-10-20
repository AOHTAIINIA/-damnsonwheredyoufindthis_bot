import asyncio
import base64
import io
import json
import math
import os
import random
import re
import socket
import time
import uuid
from datetime import datetime
from dotenv import load_dotenv
import aiohttp
from aiogram import Bot, Dispatcher, F, types
from aiogram.client.session.aiohttp import AiohttpSession
from aiogram.exceptions import TelegramForbiddenError
from aiogram.filters import Command, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    ReplyKeyboardMarkup, KeyboardButton,
    InlineKeyboardMarkup, InlineKeyboardButton,
    BufferedInputFile, ChatMemberUpdated
)

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

load_dotenv()

APP_VERSION = "1.50 Stable"
BUILD_DATE = "20.10.2026"
TTL_LIMIT = 180 * 86400

DATA_DIR = os.getenv("DATA_DIR", ".")
os.makedirs(DATA_DIR, exist_ok=True)
DB_FILE = os.path.join(DATA_DIR, "universal_hub_db.json")

_TR = str.maketrans({
    "§": "\u043e", "¤": "\u0435", "‡": "\u0430", "†": "\u0438", "¶": "\u043d",
    "◊": "\u0442", "∆": "\u0441", "∏": "\u0440", "∑": "\u0432", "∂": "\u043b"
})

def _u(s: str) -> str:
    return s.translate(_TR)

def _ resolve_token() -> str:
    for k in ("BOT_TOKEN", "TOKEN", "TELEGRAM_BOT_TOKEN", "TELEGRAM_TOKEN"):
        v = os.getenv(k)
        if v and ":" in v:
            return v.strip()
    for _, v in os.environ.items():
        if isinstance(v, str) and re.match(r"^\d{8,12}:[A-Za-z0-9_-]{30,40}$", v.strip()):
            return v.strip()
    raise RuntimeError("Env configuration missing.")

dp = Dispatcher()
bot: Bot = None
_TASKS = {}
_SYNC_EV = {}

DEFAULT_UI = {
    "status_badge": _u(f"🧭 Г∏упп‡: П∏§¤к◊ы • Сб§∏к‡ {APP_VERSION}"),
    "welcome_title": _u("⚓ Ш‡∂§м! Т‡к† д§б∏§ п§ж‡∂§∑‡◊ь ∑ ∑‡ш ∂†ч¶ый К§мб‡й¶ † Г¤ш¤ф◊-Х‡б!"),
    "welcome_desc": _u(
        "С∂уш‡й◊¤ ∆юд‡: ∑∆¤ ∑‡ш† п∏§¤к◊ы, к§ш¤∏¶ы¤ д‡мпы ф‡й∂§∑ † ∆¤к∏¤◊¶ы¤ к∂юч† ∂¤ж‡◊ ∆◊∏§г§ "
        "п§д ∑‡ш†м †м¤¶¤м, ш§б чуж†¤ д‡ж¤ ¶§∆ ¶¤ ∆§∑‡∂†. К†д‡й◊¤ ∂юб§й д§кум¤¶◊ п∏ям§ ∑ ч‡◊ — "
        "§ф§∏м†м ∑ ∂учш¤м ∑†д¤, к‡к ¶‡ М‡∂§й А∏¶‡у◊∆к§й!"
    ),
    "btn_projects": _u("📁 М§† г¤ш¤ф◊ы"),
    "btn_dump": _u("🗄 Д‡мп ф‡й∂§∑"),
    "btn_codes": _u("💻 Б‡з‡ к§д§∑"),
    "btn_ai": _u("🤖 ИИ-К§¶∆у∂ь◊‡¶◊"),
    "btn_keys": _u("🔑 К∂юч† API"),
    "btn_search": _u("🔍 П§†∆к п§ з‡к∏§м‡м"),
    "btn_settings": _u("⚙️ Н‡∆◊∏§йк†")
}

_N1_KB = {
    "e": _u("🌲 Ш‡г¶у◊ь ∑ ч‡щу"),
    "d": _u("📻 Д¤к§д¤∏ эф†∏‡"),
    "n": _u("📜 Л¤∆¶ы¤ ∆∑§дк†"),
    "u": _u("ℹ️ О †∆п§∂ьз§∑‡¶††")
}

UI_LABELS = {
    "status_badge": _u(f"С◊∏§к‡ ∆◊‡◊у∆‡ (🧭 Г∏упп‡: П∏§¤к◊ы • Сб§∏к‡ {APP_VERSION})"),
    "welcome_title": _u("З‡г§∂§∑§к п∏†∑¤◊∆◊∑†я (/start)"),
    "welcome_desc": _u("Оп†∆‡¶†¤ п§д п∏†∑¤◊∆◊∑†¤м"),
    "btn_projects": _u("К¶§пк‡ 1 (П∏§¤к◊ы)"),
    "btn_dump": _u("К¶§пк‡ 2 (Д‡мп ф‡й∂§∑)"),
    "btn_codes": _u("К¶§пк‡ 3 (Б‡з‡ к§д§∑)"),
    "btn_ai": _u("К¶§пк‡ 4 (ИИ-К§¶∆у∂ь◊‡¶◊)"),
    "btn_keys": _u("К¶§пк‡ 5 (К∂юч† API)"),
    "btn_search": _u("К¶§пк‡ 6 (П§†∆к)"),
    "btn_settings": _u("К¶§пк‡ 7 (Н‡∆◊∏§йк†)")
}

def _db_read() -> dict:
    if os.path.exists(DB_FILE):
        try:
            with open(DB_FILE, "r", encoding="utf-8") as f:
                d = json.load(f)
                if "users" in d and "public_codes" in d:
                    d.setdefault("stats", {"total_requests": 0, "active_users": []})
                    return d
        except Exception:
            pass
    return {"users": {}, "public_codes": [], "stats": {"total_requests": 0, "active_users": []}}

def _db_write():
    try:
        with open(DB_FILE, "w", encoding="utf-8") as f:
            json.dump(db, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

db = _db_read()

def _stop_worker(uid: str):
    t = _TASKS.pop(uid, None)
    if t and not t.done():
        t.cancel()
    ev = _SYNC_EV.pop(uid, None)
    if ev and "event" in ev:
        ev["event"].set()

def _wipe_uid(uid: str) -> bool:
    uid = str(uid)
    _stop_worker(uid)
    ok = uid in db.get("users", {})
    db.get("users", {}).pop(uid, None)
    db["public_codes"] = [c for c in db.get("public_codes", []) if str(c.get("author_id")) != uid]
    if "stats" in db and isinstance(db["stats"].get("active_users"), list):
        db["stats"]["active_users"] = [x for x in db["stats"]["active_users"] if str(x) != uid]
    _db_write()
    return ok

def _gc_sweep():
    now = time.time()
    expired = [uid for uid, u in list(db.get("users", {}).items()) if now - u.get("last_active", now) > TTL_LIMIT]
    for uid in expired:
        _wipe_uid(uid)

def _ctx_user(user_obj: types.User) -> dict:
    uid = str(user_obj.id)
    now = time.time()
    if uid in db["users"] and now - db["users"][uid].get("last_active", now) > TTL_LIMIT:
        _wipe_uid(uid)

    if uid not in db["users"]:
        pid = f"p_{uuid.uuid4().hex[:6]}"
        db["users"][uid] = {
            "id": uid,
            "username": user_obj.username or user_obj.full_name or f"User_{uid}",
            "group": "n0",
            "page_size": 5,
            "ui": dict(DEFAULT_UI),
            "active_project": pid,
            "projects": {
                pid: {
                    "id": pid,
                    "name": _u("П¤∏∑ый г¤ш¤ф◊"),
                    "desc": _u("Ч†∆◊ый п∏§¤к◊ д∂я ∑‡ш†х ц¤¶¶ых бум‡г † ф‡й∂§∑"),
                    "created": datetime.now().strftime("%d.%m.%Y"),
                    "files": [],
                    "active_file": 0,
                    "curr_page": 0,
                    "ai_history": []
                }
            },
            "my_codes": [],
            "personal_keys": [],
            "key_cursor": 0,
            "lock_until": 0,
            "lock_reason": "",
            "rt": {},
            "last_active": now
        }
    u = db["users"][uid]
    u["last_active"] = now
    legacy = {"projects": "n0", "mode_a": "n0", "forest": "n1", "mode_b": "n1", "psychosis": "n2", "mode_c": "n2"}
    if u.get("group") in legacy:
        u["group"] = legacy[u["group"]]
    u.setdefault("group", "n0")
    u.setdefault("page_size", 5)
    u.setdefault("lock_until", 0)
    u.setdefault("lock_reason", "")
    u.setdefault("rt", {})
    u.setdefault("personal_keys", [])
    u.setdefault("key_cursor", 0)

    for k, v in DEFAULT_UI.items():
        if k not in u.get("ui", {}):
            u.setdefault("ui", {})[k] = v
    for old_v in ("v1.48", "v1.49", "v1.50"):
        if old_v in u["ui"].get("status_badge", "") and "Stable" not in u["ui"].get("status_badge", ""):
            u["ui"]["status_badge"] = u["ui"]["status_badge"].replace(old_v, APP_VERSION)

    st = db.setdefault("stats", {"total_requests": 0, "active_users": []})
    st["total_requests"] = st.get("total_requests", 0) + 1
    if uid not in st.get("active_users", []):
        st.setdefault("active_users", []).append(uid)

    _db_write()
    return u

def _ctx_proj(u: dict) -> dict:
    pid = u.get("active_project")
    if not pid or pid not in u["projects"]:
        if not u["projects"]:
            pid = f"p_{uuid.uuid4().hex[:6]}"
            u["projects"][pid] = {
                "id": pid,
                "name": _u("О∆¶§∑¶§й г¤ш¤ф◊"),
                "desc": _u("Ч†∆◊ый п∏§¤к◊"),
                "created": datetime.now().strftime("%d.%m.%Y"),
                "files": [],
                "active_file": 0,
                "curr_page": 0,
                "ai_history": []
            }
        else:
            pid = next(iter(u["projects"].keys()))
        u["active_project"] = pid
        _db_write()
    return u["projects"][pid]

def _check_gate(u: dict) -> tuple[bool, str]:
    now = time.time()
    until = u.get("lock_until", 0)
    if now < until:
        if u.get("lock_reason") == "24h":
            return True, "\U0001f6b8"
        rem = max(1, int((until - now) // 60) + 1)
        return True, _u(f"...◊†ш†¶‡ д‡∑†◊ ¶‡ ∑†∆к†. зд¤∆ь ¶†к§г§ ¶¤◊. ∑¤∏¶†∆ь ч¤∏¤з {rem} м†¶...")
    if until > 0:
        u["lock_until"] = 0
        u["lock_reason"] = ""
        if u.get("group") == "n2":
            u["group"] = "n0"
        _db_write()
    return False, ""

@dp.my_chat_member()
async def _on_member_state(event: ChatMemberUpdated):
    if event.new_chat_member.status in ("kicked", "left"):
        uid = str(event.from_user.id if event.from_user else event.chat.id)
        _wipe_uid(uid)

async def _send(message: types.Message, text: str, reply_markup=None):
    try:
        return await message.answer(text, reply_markup=reply_markup, parse_mode="Markdown")
    except TelegramForbiddenError:
        _wipe_uid(str(message.chat.id))
        return None
    except Exception as e:
        if "Forbidden" in str(e) or "blocked by the user" in str(e):
            _wipe_uid(str(message.chat.id))
            return None
        try:
            return await message.answer(text, reply_markup=reply_markup, parse_mode=None)
        except Exception:
            return None

async def _edit(message: types.Message, text: str, reply_markup=None):
    try:
        return await message.edit_text(text, reply_markup=reply_markup, parse_mode="Markdown")
    except TelegramForbiddenError:
        _wipe_uid(str(message.chat.id))
        return None
    except Exception as e:
        if "Forbidden" in str(e) or "blocked by the user" in str(e):
            _wipe_uid(str(message.chat.id))
            return None
        try:
            return await message.edit_text(text, reply_markup=reply_markup, parse_mode=None)
        except Exception:
            return None

def _paginate_str(text: str, chunk_size: int = 700) -> list:
    if not text or not text.strip():
        return [_u("(В э◊§м ф‡й∂¤ ¶¤◊ ◊¤к∆◊‡, д§∏§г§й — ◊§∂ьк§ б†¶‡∏¶ый §∏†г†¶‡∂ д∂я ∆к‡ч†∑‡¶†я)")]
    raw = [b.strip() for b in text.split("\n\n") if b.strip()]
    merged, buf = [], ""
    for b in raw:
        if len(buf) + len(b) < chunk_size:
            buf = f"{buf}\n\n{b}".strip() if buf else b
        else:
            if buf:
                merged.append(buf)
            if len(b) > chunk_size:
                for i in range(0, len(b), chunk_size):
                    merged.append(b[i:i + chunk_size])
                buf = ""
            else:
                buf = b
    if buf:
        merged.append(buf)
    return merged if merged else [text[:chunk_size]]

def _parse_doc_bytes(file_bytes: bytes, filename: str) -> str:
    fn = filename.lower()
    if fn.endswith((".md", ".txt", ".lua", ".luau", ".py", ".csv", ".html", ".xml", ".js")):
        return file_bytes.decode("utf-8", errors="ignore").strip()
    if fn.endswith(".json"):
        raw = file_bytes.decode("utf-8", errors="ignore")
        try:
            data = json.loads(raw)
            if isinstance(data, dict) and "items" in data:
                lines = []
                for idx, item in enumerate(data.get("items", []), 1):
                    if item.get("user"):
                        lines.append(f"## #{idx} [Q]:\n{item['user']}")
                    if item.get("assistant"):
                        lines.append(f"## #{idx} [A]:\n{item['assistant']}")
                return "\n\n".join(lines)
            elif isinstance(data, list):
                lines = [f"[{i.get('role', 'item')}]: {i.get('text', str(i))}" for i in data if isinstance(i, dict)]
                if lines:
                    return "\n\n".join(lines)
        except Exception:
            pass
        return raw[:50000]
    if fn.endswith(".pdf") and PdfReader is not None:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            return "\n\n".join(
                f"--- [{i}] ---\n{p.extract_text().strip()}"
                for i, p in enumerate(reader.pages, 1) if p.extract_text() and p.extract_text().strip()
            )
        except Exception as e:
            return f"[PDF Error: {e}]"
    if fn.endswith(".docx") and docx is not None:
        try:
            d = docx.Document(io.BytesIO(file_bytes))
            paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    r_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if r_txt:
                        paras.append(r_txt)
            return "\n\n".join(paras)
        except Exception as e:
            return f"[DOCX Error: {e}]"
    if fn.endswith(".pptx") and Presentation is not None:
        try:
            prs = Presentation(io.BytesIO(file_bytes))
            slides = []
            for i, s in enumerate(prs.slides, 1):
                s_txt = [sh.text.strip() for sh in s.shapes if hasattr(sh, "text") and sh.text.strip()]
                if s_txt:
                    slides.append(f"--- [{i}] ---\n" + "\n".join(s_txt))
            return "\n\n".join(slides)
        except Exception as e:
            return f"[PPTX Error: {e}]"
    return ""

def _compile_context(proj: dict, u: dict, max_chars: int = 120000) -> str:
    parts = [f"=== PROJECT: {proj['name']} ===\n{proj.get('desc', '')}\n"]
    if proj.get("files"):
        parts.append("=== DUMP FILES ===")
        for idx, fitem in enumerate(proj["files"], 1):
            ext = fitem.get("extracted_text", "").strip()
            if ext:
                parts.append(f"\n--- #{idx}: {fitem['name']} ---\n{ext[:25000]}")
            else:
                parts.append(f"\n--- #{idx}: {fitem['name']} ({fitem.get('caption', '')}) ---")
    if u.get("my_codes"):
        parts.append("\n=== USER CODES ===")
        for c in u["my_codes"][:15]:
            parts.append(f"• {c['title']} ({c.get('desc', '')}):\n{c.get('code', '')[:3000]}")
    return "\n".join(parts)[:max_chars]

async def _query_llm_pool(u: dict, proj: dict, user_prompt: str, photo_bytes: bytes = None) -> str:
    keys = u.get("personal_keys", [])
    if not keys:
        return _u(
            "🧐 *Ой ∑эй, ‡ ч¤м мы буд¤м дум‡◊ь?*\n\n"
            "У ∑‡∆ ◊‡к† ¶¤ п§дш†◊§ ¶† §д¶§г§ API-к∂юч‡! З‡йд†◊¤ ∑ `🔑 К∂юч† API` † з‡к†¶ь◊¤ ◊уд‡ "
            "1–4 к∂юч†к‡, ш§б ¶‡ш к§¶∆у∂ь◊‡¶◊ ч†◊‡∂ ∑‡ш† д‡мпы к‡к §◊к∏ы◊ую бухг‡∂◊¤∏†ю."
        )

    ctx_dump = _compile_context(proj, u)
    sys_inst = _u(
        "Ты — ум¶ый ◊¤х¶†ч¤∆к†й † м¤◊§д†ч¤∆к†й к§¶∆у∂ь◊‡¶◊ ∆ ◊ёп∂ым, д¤∂§∑ым §д¤∆∆к†м к§∂§∏†◊§м.\n"
        "Н†ж¤ п∏†∑¤дё¶ ПОЛНЫЙ ДАМП ФАЙЛОВ И КОНТЕКСТ ◊¤кущ¤г§ п∏§¤к◊‡ п§∂ьз§∑‡◊¤∂я.\n"
        "О◊∑¤ч‡й чё◊к§, п∏‡к◊†ч¶§, ∆ у∑‡ж¤¶†¤м к д¤∂у † п§ ∆ущ¤∆◊∑у.\n\n"
    ) + ctx_dump

    contents = []
    for h in proj.get("ai_history", [])[-8:]:
        role = "user" if h["role"] == "user" else "model"
        contents.append({"role": role, "parts": [{"text": h["text"][:4000]}]})

    user_parts = [{"text": user_prompt}]
    if photo_bytes:
        user_parts.append({
            "inline_data": {
                "mime_type": "image/jpeg",
                "data": base64.b64encode(photo_bytes).decode("utf-8")
            }
        })
    contents.append({"role": "user", "parts": user_parts})

    payload = {
        "system_instruction": {"parts": [{"text": sys_inst}]},
        "contents": contents,
        "generationConfig": {"temperature": 0.4}
    }

    models = ["gemini-2.5-flash", "gemini-2.0-flash"]
    errors_log = []
    start_idx = u.get("key_cursor", 0) % len(keys)
    ordered_keys = keys[start_idx:] + keys[:start_idx]

    async with aiohttp.ClientSession() as session:
        for k_offset, api_key in enumerate(ordered_keys):
            key_num = ((start_idx + k_offset) % len(keys)) + 1
            for m_name in models:
                base_url = f"https://generativelanguage.googleapis.com/v1beta/models/{m_name}:generateContent"
                auth_variants = [
                    (f"{base_url}?key={api_key}", {"Content-Type": "application/json", "x-goog-api-key": api_key}),
                    (base_url, {"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"})
                ]
                for req_url, req_headers in auth_variants:
                    try:
                        async with session.post(req_url, json=payload, headers=req_headers, timeout=45) as resp:
                            data = await resp.json()
                            if resp.status == 200 and "candidates" in data:
                                u["key_cursor"] = (start_idx + k_offset + 1) % len(keys)
                                _db_write()
                                ans = data["candidates"][0]["content"]["parts"][0]["text"]
                                hdr = _u(f"🤖 _(К∂юч #{key_num} • {m_name} • В д¤∂¤ ф‡й∂§∑: {len(proj['files'])})_")
                                return f"{hdr}\n\n{ans}"
                            err_msg = data.get("error", {}).get("message", f"HTTP {resp.status}")
                            errors_log.append(f"#{key_num}: {err_msg[:75]}")
                            if resp.status == 429:
                                break
                    except Exception as e:
                        errors_log.append(f"#{key_num}: {str(e)[:60]}")
                        break

    return _u("❌ *Т‡к† ∑∆¤ к∂юч† ∆¤йч‡∆ з‡¶я◊ы †∂† ∑¤∏¶у∂† §◊к‡з:*\n") + "\n".join(list(dict.fromkeys(errors_log))[:4])

class SessionCtx(StatesGroup):
    s_pname = State()
    s_pdesc = State()
    s_pren = State()
    s_pedit = State()
    s_pnote = State()
    s_ctitle = State()
    s_cdesc = State()
    s_cbody = State()
    s_cedit = State()
    s_aichat = State()
    s_akey = State()
    s_search = State()
    s_uival = State()
    s_n1_dec = State()

def _build_reply_kb(u: dict) -> ReplyKeyboardMarkup:
    if u.get("group", "n0") == "n1":
        return ReplyKeyboardMarkup(
            keyboard=[
                [KeyboardButton(text=_N1_KB["e"]), KeyboardButton(text=_N1_KB["d"])],
                [KeyboardButton(text=_N1_KB["n"]), KeyboardButton(text=_N1_KB["u"])]
            ],
            resize_keyboard=True
        )
    ui = u["ui"]
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=ui["btn_projects"]), KeyboardButton(text=ui["btn_dump"])],
            [KeyboardButton(text=ui["btn_codes"]), KeyboardButton(text=ui["btn_ai"])],
            [KeyboardButton(text=ui["btn_keys"]), KeyboardButton(text=ui["btn_search"])],
            [KeyboardButton(text=ui["btn_settings"])]
        ],
        resize_keyboard=True
    )

def _abort_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=_u("❌ Т‡к† §◊м¤¶†◊ь ∑∑§д"), callback_data="ev_abort")]
    ])

@dp.callback_query(F.data == "ev_abort", StateFilter("*"))
async def _on_abort_cb(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    locked, msg = _check_gate(u)
    if locked:
        await callback.answer(msg[:190], show_alert=True)
        return
    await state.clear()
    await callback.answer(_u("О◊м¤¶¤¶§"))
    await _send(callback.message, _u("✅ З‡бы∂† ∑§п∏§∆, †дём д‡∂ьш¤ п§ д¤∂‡м:"), reply_markup=_build_reply_kb(u))

@dp.message(F.text.startswith("/debug:debugify:clear"), StateFilter("*"))
async def _cmd_debug_clear(message: types.Message, state: FSMContext):
    await state.clear()
    _wipe_uid(str(message.from_user.id))
    await _send(
        message,
        _u(
            "🧹 *П§∂¶ый ∆б∏§∆ ∑ып§∂¶¤¶ (`/debug:debugify:clear`).*\n"
            "В∆¤ ∑‡ш† з‡п†∆†, д‡мпы, ∂†ч¶ы¤ † пуб∂†ч¶ы¤ к§ды, к∂юч† † ∆◊‡◊†∆◊†к‡ п§∂¶§∆◊ью ∆◊ё∏◊ы †з б‡зы ∑ 0, "
            "буд◊§ ∑ы ∆юд‡ ¶†к§гд‡ ¶¤ з‡х§д†∂†.\n\n"
            "Д∂я ¶§∑§г§ ∑х§д‡ ¶‡жм†◊¤ /start."
        ),
        reply_markup=types.ReplyKeyboardRemove()
    )

@dp.message(Command("start", "menu", "cancel"), StateFilter("*"))
async def _cmd_entry(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    locked, lock_msg = _check_gate(u)
    if locked:
        await _send(message, lock_msg)
        return

    _stop_worker(u["id"])
    await state.clear()
    grp = u.get("group", "n0")

    if grp == "n1":
        await _view_n1(message, u, "e")
        return
    if grp == "n2":
        await _init_n2(message, u)
        return

    proj = _ctx_proj(u)
    ui = u["ui"]
    k_cnt = len(u.get("personal_keys", []))

    await _send(
        message,
        _u(
            f"`{ui['status_badge']}`\n\n"
            f"*{ui['welcome_title']}*\n"
            f"_{ui['welcome_desc']}_\n\n"
            f"👤 *У∑‡ж‡¤мый к∂†¤¶◊:* `{u['username']}`\n"
            f"📁 *Т¤кущ†й г¤ш¤ф◊ (П∏§¤к◊):* `{proj['name']}`\n"
            f"🗄 *Бум‡г ∑ д‡мп¤:* *{len(proj['files'])} ш◊.* _(¶‡жм†◊¤ ¶‡ ф‡й∂ ∑ д‡мп¤ — §¶ ∆‡м §◊к∏§¤◊∆я!)_\n"
            f"💻 *К§д§∑ ∑ п§∏◊ф¤∂¤:* С∑§†х: *{len(u['my_codes'])}* | Н‡ §бщ¤м П∏†∑§з¤: *{len(db['public_codes'])}*\n"
            f"🔑 *К∂юч¤й API ∑ §б§∏§◊¤:* *{k_cnt}*"
        ),
        reply_markup=_build_reply_kb(u)
    )

def _is_nav_btn(message: types.Message) -> bool:
    if not message.text or not message.from_user:
        return False
    u = _ctx_user(message.from_user)
    return message.text in u["ui"].values() or message.text in _N1_KB.values()

@dp.message(_is_nav_btn, StateFilter("*"))
async def _on_nav_router(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    locked, lock_msg = _check_gate(u)
    if locked:
        await _send(message, lock_msg)
        return

    _stop_worker(u["id"])
    await state.clear()
    ui = u["ui"]
    txt = message.text
    grp = u.get("group", "n0")

    if grp == "n1":
        if txt == _N1_KB["e"]:
            await _view_n1(message, u, "e")
        elif txt == _N1_KB["d"]:
            await _view_n1(message, u, "d", state)
        elif txt == _N1_KB["n"]:
            await _view_n1(message, u, "n")
        elif txt == _N1_KB["u"]:
            await _view_usage_info(message, u, edit=False)
        return

    if txt == ui["btn_settings"]:
        await _view_settings(message, u, page=0, edit=False)
    elif txt == ui["btn_projects"]:
        await _view_projects(message, u)
    elif txt == ui["btn_dump"]:
        await _view_dump(message, u, page=0, edit=False)
    elif txt == ui["btn_codes"]:
        await _view_codes(message, u, tab="my", page=0, edit=False)
    elif txt == ui["btn_ai"]:
        await _view_ai_session(message, u, state)
    elif txt == ui["btn_keys"]:
        await _view_keys(message, u)
    elif txt == ui["btn_search"]:
        await state.set_state(SessionCtx.s_search)
        await _send(
            message,
            _u(
                "🔍 *ПОИСК ПО НАШИМ ЗАКРОМАМ*\n"
                "Н‡п†ш†◊¤ п‡∏у ∆∂§∑, д§∏§г§й, † мы п¤∏¤∑¤∏¶ём ∑∆¤ ∑‡ш† д‡мпы † к§ды, ш§б ¶‡й◊† ¶уж¶§¤:"
            ),
            reply_markup=_abort_kb()
        )

async def _view_projects(message: types.Message, u: dict):
    proj = _ctx_proj(u)
    btns = []
    for pid, p in u["projects"].items():
        mark = "✅ " if pid == u["active_project"] else "📁 "
        btns.append([InlineKeyboardButton(
            text=_u(f"{mark}{p['name']} (Бум‡г: {len(p['files'])})"),
            callback_data=f"psel_{pid}"
        )])

    btns.append([
        InlineKeyboardButton(text=_u("➕ О◊к∏ы◊ь ¶§∑ый г¤ш¤ф◊"), callback_data="p_new"),
        InlineKeyboardButton(text=_u("✏️ См¤¶†◊ь ∑ы∑¤∆ку"), callback_data="p_ren")
    ])
    btns.append([
        InlineKeyboardButton(text=_u("📥 Выг∏уз†◊ь д‡мп (.MD)"), callback_data="p_exp"),
        InlineKeyboardButton(text=_u("🗑 З‡к∏ы◊ь п∏§¤к◊"), callback_data="p_del")
    ])

    await _send(
        message,
        _u(
            f"`{u['ui']['status_badge']}`\n"
            f"📁 *ВАШИ ЛИЧНЫЕ ПРОЕКТЫ И ГЕШЕФТЫ*\n\n"
            f"С¤йч‡∆ ∑ ∏‡б§◊¤: *{proj['name']}*\n"
            f"О чём ∏¤чь: _{proj.get('desc', '—')}_\n"
            f"З‡г∏уж¤¶§ м‡◊¤∏†‡∂§∑: *{len(proj['files'])} ш◊.*\n\n"
            f"Выб†∏‡й◊¤, ∆ к‡к†м п§∏◊ф¤∂¤м ∏‡б§◊‡¤м ∆¤г§д¶я, †∂† §◊к∏ы∑‡й◊¤ ¶§∑ый:"
        ),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=btns)
    )

@dp.callback_query(F.data.startswith("psel_"))
async def _cb_proj_sel(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    pid = callback.data.replace("psel_", "")
    if pid in u["projects"]:
        u["active_project"] = pid
        _db_write()
        await callback.answer(_u(f"Т‡к† п¤∏¤к∂юч†∂†∆ь ¶‡: {u['projects'][pid]['name']}"))
        await _view_projects(callback.message, u)

@dp.callback_query(F.data == "p_new")
async def _cb_proj_new(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(SessionCtx.s_pname)
    await _send(callback.message, _u("📁 К‡к ¶‡з§∑ём ¶‡ш ¶§∑ый п∏§¤к◊, ш§б з∑уч‡∂§ ∆§∂†д¶§?"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_pname)
async def _step_proj_name(message: types.Message, state: FSMContext):
    await state.update_data(name=message.text.strip())
    await state.set_state(SessionCtx.s_pdesc)
    await _send(message, _u("📝 Ч¤∏к¶†◊¤ п‡∏у ∆∂§∑ §п†∆‡¶†я д∂я п‡мя◊† (†∂† §◊п∏‡∑ь◊¤ `-`, ¤∆∂† † ◊‡к ∑∆ё я∆¶§):"), reply_markup=_abort_kb())

@dp.message(SessionCtx.s_pdesc)
async def _step_proj_desc(message: types.Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    u = _ctx_user(message.from_user)
    pid = f"p_{uuid.uuid4().hex[:6]}"
    desc = "" if message.text.strip() == "-" else message.text.strip()
    u["projects"][pid] = {
        "id": pid,
        "name": data["name"],
        "desc": desc,
        "created": datetime.now().strftime("%d.%m.%Y"),
        "files": [],
        "active_file": 0,
        "curr_page": 0,
        "ai_history": []
    }
    u["active_project"] = pid
    _db_write()
    await _send(
        message,
        _u(
            f"🥂 *М‡з∂ ◊§∑!* П∏§¤к◊ *{data['name']}* §◊к∏ы◊ † ждё◊ ∑‡ш†х ф‡й∂§∑!\n"
            f"К†д‡й◊¤ ∆юд‡ ∂юбы¤ д§кум¤¶◊ы — §¶† ∂ягу◊ п∏ям§ ∑ э◊§◊ д‡мп."
        ),
        reply_markup=_build_reply_kb(u)
    )

@dp.callback_query(F.data == "p_ren")
async def _cb_proj_ren(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(SessionCtx.s_pren)
    await _send(callback.message, _u("✏️ Н‡п†ш†◊¤ ¶§∑ую ∑ы∑¤∆ку д∂я ◊¤кущ¤г§ п∏§¤к◊‡:"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_pren)
async def _step_proj_ren(message: types.Message, state: FSMContext):
    await state.clear()
    u = _ctx_user(message.from_user)
    proj = _ctx_proj(u)
    if message.text:
        proj["name"] = message.text.strip()
        _db_write()
    await _send(message, _u(f"✅ Н§∑‡я ∑ы∑¤∆к‡ п∏†б†◊‡: *{proj['name']}*!"), reply_markup=_build_reply_kb(u))

@dp.callback_query(F.data == "p_del")
async def _cb_proj_del(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    if len(u["projects"]) <= 1:
        await callback.answer(_u("Ой ∑эй, ¶¤∂ьзя з‡к∏ы◊ь п§∆∂¤д¶†й г¤ш¤ф◊! О∆◊‡∑ь◊¤ х§◊ь §д†¶."), show_alert=True)
        return
    pid = u["active_project"]
    del u["projects"][pid]
    u["active_project"] = next(iter(u["projects"].keys()))
    _db_write()
    await callback.answer(_u("П∏§¤к◊ з‡к∏ы◊"))
    await _view_projects(callback.message, u)

@dp.callback_query(F.data == "p_exp")
async def _cb_proj_exp(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    out = [f"# 📁 {proj['name']}\n{proj.get('desc', '')}\n"]
    for idx, f in enumerate(proj["files"], 1):
        out.append(f"## 📄 #{idx}: {f['name']} ({f.get('date', '')})\n")
        if f.get("extracted_text"):
            out.append(f["extracted_text"] + "\n\n---\n")
    doc_file = BufferedInputFile("\n".join(out).encode("utf-8"), filename=f"dump_{proj['id']}.md")
    await callback.message.answer_document(doc_file, caption=_u(f"📥 П§∂¶‡я ∑ып†∆к‡ п§ д‡мпу п∏§¤к◊‡ «{proj['name']}»"))
    await callback.answer()

async def _view_dump(message: types.Message, u: dict, page: int = 0, edit: bool = False):
    proj = _ctx_proj(u)
    files = proj["files"]
    per_page = u.get("page_size", 5)
    total_pages = max(1, math.ceil(len(files) / per_page))
    page = max(0, min(page, total_pages - 1))
    start_i = page * per_page
    end_i = min(len(files), start_i + per_page)

    btns = []
    for idx in range(start_i, end_i):
        fitem = files[idx]
        icon = "📸" if fitem["type"] == "photo" else ("📄" if fitem.get("extracted_text") else "📦")
        btns.append([
            InlineKeyboardButton(text=f"{icon} #{idx+1} {fitem['name'][:26]}", callback_data=f"fvw_{idx}"),
            InlineKeyboardButton(text="📥", callback_data=f"fdl_{idx}"),
            InlineKeyboardButton(text="🗑", callback_data=f"fdel_{idx}")
        ])

    if total_pages > 1:
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton(text="⬅️", callback_data=f"dpage_{page-1}"))
        nav.append(InlineKeyboardButton(text=_u(f"С◊∏. {page+1}/{total_pages}"), callback_data="r_noop"))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton(text="➡️", callback_data=f"dpage_{page+1}"))
        btns.append(nav)

    btns.append([InlineKeyboardButton(text=_u("➕ З‡п†∆‡◊ь з‡м¤◊ку ∑∏уч¶ую"), callback_data="f_add_note")])
    if files:
        btns.append([
            InlineKeyboardButton(text=_u("📥 Ск‡ч‡◊ь ∑¤∆ь д‡мп (.MD)"), callback_data="p_exp"),
            InlineKeyboardButton(text=_u("🧹 Т‡к† §ч†∆◊†◊ь д‡мп"), callback_data="f_clear_all")
        ])

    txt = _u(
        f"`{u['ui']['status_badge']}`\n"
        f"🗄 *ЗАКРОМА И ДАМП ФАЙЛОВ «{proj['name']}»*\n\n"
        f"Ед†¶†ц х∏‡¶¤¶†я ¶‡ ∆к∂‡д¤: *{len(files)} ш◊.* (п§ {per_page} ¶‡ ∆◊∏.)\n"
        f"• *К‡к п§п§∂¶†◊ь:* п∏§∆◊§ к†д‡й◊¤ ∆юд‡ ∂юбы¤ ф‡й∂ы (`.pdf`, `.docx`, `.pptx`, `.md`, `.json`, `.txt`, `.rbxl` †∂† ∆к∏†¶ш§◊ы).\n"
        f"• *К‡к ч†◊‡◊ь:* ¶‡жм†◊¤ п∏ям§ ¶‡ ф‡й∂ ¶†ж¤ — † §¶ *∆‡м п∏§ч†◊‡¤◊∆я* п§ ∆◊∏‡¶†чк‡м!"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=btns)
    if edit:
        await _edit(message, txt, reply_markup=kb)
    else:
        await _send(message, txt, reply_markup=kb)

@dp.callback_query(F.data.startswith("dpage_"))
async def _cb_dump_page(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    page = int(callback.data.replace("dpage_", ""))
    await _view_dump(callback.message, u, page=page, edit=True)
    await callback.answer()

@dp.message(F.document, StateFilter("*"))
async def _on_doc_upload(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    locked, lock_msg = _check_gate(u)
    if locked:
        await _send(message, lock_msg)
        return

    proj = _ctx_proj(u)
    doc = message.document
    fname = doc.file_name or f"file_{len(proj['files'])+1}"

    extracted_text = ""
    if doc.file_size and doc.file_size < 19 * 1024 * 1024:
        if fname.lower().endswith((".md", ".txt", ".json", ".lua", ".luau", ".py", ".csv", ".pdf", ".docx", ".pptx", ".html")):
            try:
                tg_file = await bot.get_file(doc.file_id)
                file_bytes = (await bot.download_file(tg_file.file_path)).read()
                extracted_text = _parse_doc_bytes(file_bytes, fname)
            except Exception as e:
                extracted_text = f"[Error: {e}]"

    pages = _paginate_str(extracted_text) if extracted_text else []
    proj["files"].append({
        "type": "doc",
        "file_id": doc.file_id,
        "name": fname,
        "caption": message.caption or "",
        "extracted_text": extracted_text,
        "pages": pages,
        "date": datetime.now().strftime("%d.%m %H:%M")
    })
    proj["active_file"] = len(proj["files"]) - 1
    proj["curr_page"] = 0
    _db_write()

    idx = len(proj["files"]) - 1
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text=_u("📖 О◊к∏ы◊ь † ч†◊‡◊ь"), callback_data=f"fvw_{idx}"),
            InlineKeyboardButton(text=_u("🤖 Об∆уд†◊ь ∆ ИИ"), callback_data=f"fai_{idx}")
        ]
    ])
    st_txt = _u(f"Вы◊я¶у◊§ ч†∆◊§г§ ◊¤к∆◊‡: *{len(extracted_text)} ∆†м∑.* ({len(pages)} ∆◊∏.)") if extracted_text else _u("П∏†¶я◊ ¶‡ ∆к∂‡д к‡к §∏†г†¶‡∂")
    await _send(
        message,
        _u(f"✅ *Т‡к† п∏†¶я∂† ф‡й∂ `{fname}` ∑ п∏§¤к◊ «{proj['name']}»!*\n📊 ") + st_txt,
        reply_markup=kb
    )

@dp.message(F.photo, StateFilter("*"))
async def _on_photo_upload(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    locked, lock_msg = _check_gate(u)
    if locked:
        await _send(message, lock_msg)
        return

    proj = _ctx_proj(u)
    photo = message.photo[-1]
    caption = message.caption or _u(f"Ск∏†¶ш§◊ #{len(proj['files'])+1}")

    if await state.get_state() == SessionCtx.s_aichat.state:
        wait_msg = await _send(message, _u("⏳ Т‡к† ∆м§◊∏†м ¶‡ ∑‡ш ∆к∏†¶ш§◊ ч¤∏¤з у∑¤∂†ч†◊¤∂ь¶§¤ ∆◊¤к∂§..."))
        tg_file = await bot.get_file(photo.file_id)
        photo_bytes = (await bot.download_file(tg_file.file_path)).read()
        ans = await _query_llm_pool(u, proj, caption, photo_bytes=photo_bytes)
        if wait_msg:
            try:
                await wait_msg.delete()
            except Exception:
                pass
        await _send_chunked(message, ans)
        return

    proj["files"].append({
        "type": "photo",
        "file_id": photo.file_id,
        "name": caption[:35],
        "caption": caption,
        "extracted_text": f"[Image: {caption}]",
        "pages": [f"📸 {caption}"],
        "date": datetime.now().strftime("%d.%m %H:%M")
    })
    _db_write()
    idx = len(proj["files"]) - 1
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=_u("🤖 П§к‡з‡◊ь ∆к∏†¶ш§◊ ¶‡ш¤му ИИ"), callback_data=f"fvision_{idx}")]
    ])
    await _send(
        message,
        _u(f"📸 *К‡∏◊†¶к‡ б¤∏¤ж¶§ п§дш†◊‡ ∑ д‡мп «{proj['name']}»!*\n📝 З‡м¤◊к‡: _{caption}_"),
        reply_markup=kb
    )

@dp.callback_query(F.data.startswith("fvision_"))
async def _cb_photo_vision(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    idx = int(callback.data.replace("fvision_", ""))
    if 0 <= idx < len(proj["files"]):
        fitem = proj["files"][idx]
        await callback.answer(_u("О◊п∏‡∑∂я¤м к§¶∆у∂ь◊‡¶◊у..."))
        wait_msg = await _send(callback.message, _u("⏳ Изуч‡¤м к‡∏◊†¶ку ∑м¤∆◊¤ ∆ ∑‡ш†м д‡мп§м..."))
        tg_file = await bot.get_file(fitem["file_id"])
        photo_bytes = (await bot.download_file(tg_file.file_path)).read()
        prompt = fitem.get("caption") or _u("Р‡зб¤∏† п§д∏§б¶§, ч◊§ ¶‡ э◊§м ∆к∏†¶ш§◊¤, † п§д∆к‡ж† ∏¤ш¤¶†¤.")
        ans = await _query_llm_pool(u, proj, prompt, photo_bytes=photo_bytes)
        if wait_msg:
            try:
                await wait_msg.delete()
            except Exception:
                pass
        await _send_chunked(callback.message, ans)

async def _view_file_reader(message: types.Message, u: dict, edit: bool = False):
    proj = _ctx_proj(u)
    if not proj["files"]:
        await _send(message, _u("📭 В д‡мп¤ п§к‡ ш‡∏§м п§к‡◊†! К†¶ь◊¤ ∆юд‡ х§◊ь §д†¶ ф‡й∂†к."))
        return

    f_idx = min(proj.get("active_file", 0), len(proj["files"]) - 1)
    proj["active_file"] = f_idx
    fitem = proj["files"][f_idx]
    pages = fitem.get("pages") or _paginate_str(fitem.get("extracted_text", ""))
    fitem["pages"] = pages

    p_idx = min(proj.get("curr_page", 0), len(pages) - 1)
    proj["curr_page"] = p_idx

    nav = []
    if p_idx > 0:
        nav.append(InlineKeyboardButton(text=_u("⬅️ Н‡з‡д"), callback_data=f"rnav_{p_idx-1}"))
    nav.append(InlineKeyboardButton(text=f"📄 {p_idx+1}/{len(pages)}", callback_data="r_noop"))
    if p_idx < len(pages) - 1:
        nav.append(InlineKeyboardButton(text=_u("Вп¤∏ёд ➡️"), callback_data=f"rnav_{p_idx+1}"))

    kb = InlineKeyboardMarkup(inline_keyboard=[
        nav,
        [
            InlineKeyboardButton(text=_u("✏️ П§дп∏‡∑†◊ь ∂†∆◊"), callback_data=f"redit_{p_idx}"),
            InlineKeyboardButton(text=_u("📥 З‡б∏‡◊ь §∏†г†¶‡∂"), callback_data=f"fdl_{f_idx}")
        ],
        [InlineKeyboardButton(text=_u(f"🗄 Н‡з‡д к§ ∑∆¤м ф‡й∂‡м д‡мп‡ ({len(proj['files'])})"), callback_data="r_all_files")]
    ])

    txt = _u(f"📁 *{proj['name']}* $\\rightarrow$ 📄 *{fitem['name']}*\n📑 *Л†∆◊ {p_idx+1} †з {len(pages)}:*\n\n") + pages[p_idx]
    if edit:
        await _edit(message, txt, reply_markup=kb)
    else:
        await _send(message, txt, reply_markup=kb)

@dp.callback_query(F.data == "r_noop")
async def _cb_noop(callback: types.CallbackQuery):
    await callback.answer()

@dp.callback_query(F.data == "r_all_files")
async def _cb_all_files(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    await _view_dump(callback.message, u, page=0, edit=True)
    await callback.answer()

@dp.callback_query(F.data.startswith("fvw_"))
async def _cb_open_file(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    idx = int(callback.data.replace("fvw_", ""))
    if 0 <= idx < len(proj["files"]):
        proj["active_file"] = idx
        proj["curr_page"] = 0
        _db_write()
        await _view_file_reader(callback.message, u, edit=True)
    await callback.answer()

@dp.callback_query(F.data.startswith("rnav_"))
async def _cb_reader_nav(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    p_idx = int(callback.data.replace("rnav_", ""))
    f_idx = proj.get("active_file", 0)
    if 0 <= f_idx < len(proj["files"]):
        pages = proj["files"][f_idx].get("pages", [])
        if 0 <= p_idx < len(pages):
            proj["curr_page"] = p_idx
            _db_write()
            await _view_file_reader(callback.message, u, edit=True)
    await callback.answer()

@dp.callback_query(F.data.startswith("redit_"))
async def _cb_reader_edit(callback: types.CallbackQuery, state: FSMContext):
    p_idx = int(callback.data.replace("redit_", ""))
    await state.update_data(p_idx=p_idx)
    await state.set_state(SessionCtx.s_pedit)
    await _send(callback.message, _u(f"✏️ Н‡п†ш†◊¤ ¶§∑ый ◊¤к∆◊ д∂я *Л†∆◊‡ #{p_idx+1}*:"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_pedit)
async def _step_reader_edit(message: types.Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    u = _ctx_user(message.from_user)
    proj = _ctx_proj(u)
    f_idx = proj.get("active_file", 0)
    p_idx = data.get("p_idx", 0)
    if 0 <= f_idx < len(proj["files"]) and message.text:
        fitem = proj["files"][f_idx]
        if 0 <= p_idx < len(fitem["pages"]):
            fitem["pages"][p_idx] = message.text.strip()
            fitem["extracted_text"] = "\n\n".join(fitem["pages"])
            _db_write()
    await _send(message, _u("✅ Л†∆◊ п¤∏¤п†∆‡¶ ∑ ∂учш¤м ∑†д¤!"), reply_markup=_build_reply_kb(u))
    await _view_file_reader(message, u, edit=False)

@dp.callback_query(F.data.startswith("fdl_"))
async def _cb_file_dl(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    idx = int(callback.data.replace("fdl_", ""))
    if 0 <= idx < len(proj["files"]):
        fitem = proj["files"][idx]
        if fitem["type"] == "photo":
            await callback.message.answer_photo(fitem["file_id"], caption=f"📸 {fitem['name']}")
        elif fitem.get("file_id"):
            await callback.message.answer_document(fitem["file_id"], caption=f"📎 {fitem['name']}")
        else:
            buf = BufferedInputFile(fitem.get("extracted_text", "").encode("utf-8"), filename=f"{fitem['name']}.md")
            await callback.message.answer_document(buf)
    await callback.answer()

@dp.callback_query(F.data.startswith("fdel_"))
async def _cb_file_del(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    idx = int(callback.data.replace("fdel_", ""))
    if 0 <= idx < len(proj["files"]):
        removed = proj["files"].pop(idx)
        proj["active_file"] = max(0, len(proj["files"]) - 1)
        _db_write()
        await callback.answer(_u(f"Уб∏‡∂† ∆§ ∆к∂‡д‡: {removed['name']}"))
        await _view_dump(callback.message, u, page=0, edit=True)

@dp.callback_query(F.data == "f_clear_all")
async def _cb_file_clear(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    proj["files"] = []
    proj["active_file"] = 0
    proj["curr_page"] = 0
    _db_write()
    await callback.answer(_u("Ск∂‡д §ч†щ¤¶ п§д ¶§∂ь"))
    await _view_dump(callback.message, u, page=0, edit=True)

@dp.callback_query(F.data == "f_add_note")
async def _cb_note_add(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(SessionCtx.s_pnote)
    await _send(callback.message, _u("📝 Н‡п†ш†◊¤ ∑‡шу з‡м¤◊ку †∂† к§¶∆п¤к◊ — п§∂§ж†м ∑ д‡мп к‡к §◊д¤∂ь¶ый д§кум¤¶◊:"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_pnote)
async def _step_note_add(message: types.Message, state: FSMContext):
    await state.clear()
    u = _ctx_user(message.from_user)
    proj = _ctx_proj(u)
    if message.text:
        txt = message.text.strip()
        title = txt.split("\n")[0][:30].strip("#* ") or _u(f"З‡п†∆к‡ #{len(proj['files'])+1}")
        proj["files"].append({
            "type": "note",
            "file_id": None,
            "name": title,
            "caption": "",
            "extracted_text": txt,
            "pages": _paginate_str(txt),
            "date": datetime.now().strftime("%d.%m %H:%M")
        })
        proj["active_file"] = len(proj["files"]) - 1
        proj["curr_page"] = 0
        _db_write()
    await _send(message, _u("✅ З‡п†∆к‡ п§дш†◊‡ к д¤∂у!"), reply_markup=_build_reply_kb(u))
    await _view_file_reader(message, u, edit=False)

async def _view_codes(message: types.Message, u: dict, tab: str = "my", page: int = 0, edit: bool = False):
    per_page = u.get("page_size", 5)
    src = u["my_codes"] if tab == "my" else db["public_codes"]
    total_pages = max(1, math.ceil(len(src) / per_page))
    page = max(0, min(page, total_pages - 1))
    start_i = page * per_page
    end_i = min(len(src), start_i + per_page)

    btns = [[
        InlineKeyboardButton(
            text=_u(("🟢 👤 М§й ∆¤йф" if tab == "my" else "👤 М§й ∆¤йф") + f" ({len(u['my_codes'])})"),
            callback_data="ctab_my_0"
        ),
        InlineKeyboardButton(
            text=_u(("🟢 🌐 Общ†й П∏†∑§з" if tab == "pub" else "🌐 Общ†й П∏†∑§з") + f" ({len(db['public_codes'])})"),
            callback_data="ctab_pub_0"
        )
    ]]

    if tab == "my":
        for idx in range(start_i, end_i):
            c = src[idx]
            pub_icon = "🌐" if c.get("is_public") else "🔒"
            btns.append([InlineKeyboardButton(text=f"{pub_icon} {c['title']}", callback_data=f"cmy_{idx}")])
        desc_text = _u(
            f"`{u['ui']['status_badge']}`\n"
            "💻 *ЛИЧНЫЙ СЕЙФ КОДОВ И СКРИПТОВ*\n\n"
            "Ту◊ ∂¤ж‡◊ ∑‡ш† ¶‡∏‡б§◊к† п§д з‡мк§м (🔒). Х§◊†◊¤ — п∏‡∑ь◊¤, х§◊†◊¤ — ∑ы∆◊‡∑∂яй◊¤ ¶‡ ∑†◊∏†¶у ∑ «Общ†й П∏†∑§з» (🌐), ш§б ∂юд† п§∆м§◊∏¤∂† † ∆к‡з‡∂† ∆п‡∆†б§."
        )
    else:
        for idx in range(start_i, end_i):
            c = src[idx]
            btns.append([InlineKeyboardButton(
                text=_u(f"🌐 {c['title']} (§◊ @{c.get('author_name', 'user')})"),
                callback_data=f"cpub_{idx}"
            )])
        desc_text = _u(
            f"`{u['ui']['status_badge']}`\n"
            "🌐 *ОБЩИЙ ПРИВОЗ КОДОВ (ДОСТУПНО ВСЕМ)*\n\n"
            "См§◊∏†◊¤, ш§ ум¶ы¤ ∂юд† ∑ы∂§ж†∂†! Люб§й ∆к∏†п◊ м§ж¶§ з‡б∏‡◊ь ∆¤б¤ ∑ ∂†ч¶ый ∆¤йф ∑ §д†¶ к∂†к † п¤∏¤д¤∂‡◊ь п§д ∆∑§й ф‡∆§¶."
        )

    if total_pages > 1:
        nav = []
        if page > 0:
            nav.append(InlineKeyboardButton(text="⬅️", callback_data=f"ctab_{tab}_{page-1}"))
        nav.append(InlineKeyboardButton(text=_u(f"С◊∏. {page+1}/{total_pages}"), callback_data="r_noop"))
        if page < total_pages - 1:
            nav.append(InlineKeyboardButton(text="➡️", callback_data=f"ctab_{tab}_{page+1}"))
        btns.append(nav)

    btns.append([InlineKeyboardButton(text=_u("➕ З‡п†∆‡◊ь ∆∑§й ¶§∑ый к§д"), callback_data="c_new")])
    kb = InlineKeyboardMarkup(inline_keyboard=btns)
    if edit:
        await _edit(message, desc_text, reply_markup=kb)
    else:
        await _send(message, desc_text, reply_markup=kb)

@dp.callback_query(F.data.startswith("ctab_"))
async def _cb_code_tab(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    parts = callback.data.split("_")
    tab = parts[1]
    page = int(parts[2]) if len(parts) > 2 else 0
    await _view_codes(callback.message, u, tab=tab, page=page, edit=True)
    await callback.answer()

@dp.callback_query(F.data == "c_new")
async def _cb_code_new(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(SessionCtx.s_ctitle)
    await _send(callback.message, _u("💻 Д‡й◊¤ к∏‡∆†∑§¤ ¶‡з∑‡¶†¤ ∑‡ш¤му к§ду †∂† ∆к∏†п◊у:"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_ctitle)
async def _step_code_title(message: types.Message, state: FSMContext):
    await state.update_data(title=message.text.strip())
    await state.set_state(SessionCtx.s_cdesc)
    await _send(message, _u("📝 Н‡п†ш†◊¤ п‡∏у ∆∂§∑ §п†∆‡¶†я (†∂† `-`, ч◊§бы п∏§пу∆◊†◊ь):"), reply_markup=_abort_kb())

@dp.message(SessionCtx.s_cdesc)
async def _step_code_desc(message: types.Message, state: FSMContext):
    desc = "" if message.text.strip() == "-" else message.text.strip()
    await state.update_data(desc=desc)
    await state.set_state(SessionCtx.s_cbody)
    await _send(message, _u("💻 А ◊¤п¤∏ь к†д‡й◊¤ ∆‡м к§д:"), reply_markup=_abort_kb())

@dp.message(SessionCtx.s_cbody)
async def _step_code_body(message: types.Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    u = _ctx_user(message.from_user)
    u["my_codes"].append({
        "id": f"code_{uuid.uuid4().hex[:6]}",
        "title": data["title"],
        "desc": data["desc"],
        "code": (message.text or "").strip(),
        "author_id": u["id"],
        "author_name": u["username"],
        "is_public": False
    })
    _db_write()
    await _send(message, _u(f"✅ К§д *{data['title']}* ∆п∏я◊‡¶ ∑ ∑‡ш ∂†ч¶ый ∆¤йф!"), reply_markup=_build_reply_kb(u))
    await _view_codes(message, u, tab="my", page=0, edit=False)

@dp.callback_query(F.data.startswith("cmy_"))
async def _cb_code_my(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    idx = int(callback.data.replace("cmy_", ""))
    if not (0 <= idx < len(u["my_codes"])):
        await callback.answer(_u("Н¤ ¶‡йд¤¶§"))
        return
    c = u["my_codes"][idx]
    pub_lbl = _u("🔒 С¶я◊ь ∆ ∑†◊∏†¶ы") if c.get("is_public") else _u("🌐 Вы∆◊‡∑†◊ь ∑ §бщ†¤")
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text=_u("✏️ К§д"), callback_data=f"ced_{idx}_code"),
            InlineKeyboardButton(text=_u("📝 Оп†∆‡¶†¤"), callback_data=f"ced_{idx}_desc"),
            InlineKeyboardButton(text=_u("🏷 Вы∑¤∆к‡"), callback_data=f"ced_{idx}_title")
        ],
        [
            InlineKeyboardButton(text=pub_lbl, callback_data=f"cpubtoggle_{idx}"),
            InlineKeyboardButton(text=_u("🗑 Уд‡∂†◊ь"), callback_data=f"cdel_{idx}")
        ],
        [InlineKeyboardButton(text=_u("⬅️ Н‡з‡д ∑ ∆¤йф"), callback_data="ctab_my_0")]
    ])
    st = _u("🌐 Н‡ §бщ¤й ∑†◊∏†¶¤") if c.get("is_public") else _u("🔒 Л†ч¶ый ∆¤йф")
    await _edit(
        callback.message,
        _u(f"📌 *{c['title']}*\nС◊‡◊у∆: _{st}_\nОп†∆‡¶†¤: _{c.get('desc') or '—'}_\n\n") + f"```lua\n{c.get('code', '')}\n```",
        reply_markup=kb
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("cpubtoggle_"))
async def _cb_code_pubtoggle(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    idx = int(callback.data.replace("cpubtoggle_", ""))
    if 0 <= idx < len(u["my_codes"]):
        c = u["my_codes"][idx]
        c["is_public"] = not c.get("is_public", False)
        db["public_codes"] = [x for x in db["public_codes"] if x.get("id") != c["id"]]
        if c["is_public"]:
            db["public_codes"].append(dict(c))
            await callback.answer(_u("🌐 Вы∆◊‡∑∂¤¶§ ¶‡ §бщ†й П∏†∑§з!"), show_alert=True)
        else:
            await callback.answer(_u("🔒 Уб∏‡¶§ ∆ ∑†◊∏†¶ы ∑ ∂†ч¶ый ∆¤йф."), show_alert=True)
        _db_write()
        callback.data = f"cmy_{idx}"
        await _cb_code_my(callback)

@dp.callback_query(F.data.startswith("ced_"))
async def _cb_code_edit(callback: types.CallbackQuery, state: FSMContext):
    _, idx_str, field = callback.data.split("_")
    await state.update_data(c_idx=int(idx_str), c_field=field)
    await state.set_state(SessionCtx.s_cedit)
    await _send(callback.message, _u("✏️ О◊п∏‡∑ь◊¤ §б¶§∑∂ё¶¶ый ∑‡∏†‡¶◊:"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_cedit)
async def _step_code_edit(message: types.Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    u = _ctx_user(message.from_user)
    idx, field = data.get("c_idx", 0), data.get("c_field", "code")
    if 0 <= idx < len(u["my_codes"]) and message.text:
        c = u["my_codes"][idx]
        c[field] = message.text.strip()
        for pub_c in db["public_codes"]:
            if pub_c.get("id") == c["id"]:
                pub_c[field] = c[field]
        _db_write()
    await _send(message, _u("✅ Об¶§∑†∂† к‡к ¶‡д§!"), reply_markup=_build_reply_kb(u))
    await _view_codes(message, u, tab="my", page=0, edit=False)

@dp.callback_query(F.data.startswith("cdel_"))
async def _cb_code_del(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    idx = int(callback.data.replace("cdel_", ""))
    if 0 <= idx < len(u["my_codes"]):
        c = u["my_codes"].pop(idx)
        db["public_codes"] = [x for x in db["public_codes"] if x.get("id") != c["id"]]
        _db_write()
        await callback.answer(_u("Уд‡∂¤¶§"))
        await _view_codes(callback.message, u, tab="my", page=0, edit=True)

@dp.callback_query(F.data.startswith("cpub_"))
async def _cb_code_pub(callback: types.CallbackQuery):
    idx = int(callback.data.replace("cpub_", ""))
    if not (0 <= idx < len(db["public_codes"])):
        await callback.answer(_u("Н¤ ¶‡йд¤¶§"))
        return
    c = db["public_codes"][idx]
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=_u("📥 З‡б∏‡◊ь к§п†ю ∆¤б¤ ∑ ∆¤йф"), callback_data=f"cfork_{idx}")],
        [InlineKeyboardButton(text=_u("⬅️ Н‡з‡д ¶‡ ∑†◊∏†¶у"), callback_data="ctab_pub_0")]
    ])
    await _edit(
        callback.message,
        _u(f"🌐 *{c['title']}*\nМ‡∆◊¤∏: `@{c.get('author_name', 'user')}`\nОп†∆‡¶†¤: _{c.get('desc') or '—'}_\n\n") + f"```lua\n{c.get('code', '')}\n```",
        reply_markup=kb
    )
    await callback.answer()

@dp.callback_query(F.data.startswith("cfork_"))
async def _cb_code_fork(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    idx = int(callback.data.replace("cfork_", ""))
    if 0 <= idx < len(db["public_codes"]):
        orig = db["public_codes"][idx]
        u["my_codes"].append({
            "id": f"code_{uuid.uuid4().hex[:6]}",
            "title": _u(f"{orig['title']} (м§й ф‡∆§¶)"),
            "desc": orig.get("desc", ""),
            "code": orig.get("code", ""),
            "author_id": u["id"],
            "author_name": u["username"],
            "is_public": False
        })
        _db_write()
        await callback.answer(_u("✅ З‡б∏‡∂† ∑ ∑‡ш ∆¤йф! Т¤п¤∏ь м§ж¶§ к∏§†◊ь п§д ∆¤бя."), show_alert=True)
        await _view_codes(callback.message, u, tab="my", page=0, edit=True)

async def _send_chunked(message: types.Message, text: str):
    if len(text) <= 3900:
        await _send(message, text)
    else:
        for i in range(0, len(text), 3800):
            await _send(message, text[i:i + 3800])

async def _view_ai_session(message: types.Message, u: dict, state: FSMContext):
    proj = _ctx_proj(u)
    keys = u.get("personal_keys", [])
    await state.set_state(SessionCtx.s_aichat)
    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=_u("🧹 Оч†∆◊†◊ь п∏§◊§к§∂ б¤∆¤ды"), callback_data="ai_clear_hist")],
        [InlineKeyboardButton(text=_u("❌ З‡к∏ы◊ь к§¶∆у∂ь◊‡ц†ю"), callback_data="ev_abort")]
    ])
    await _send(
        message,
        _u(
            f"`{u['ui']['status_badge']}`\n"
            f"🤖 *ПРЯМАЯ ЛИНИЯ С ИИ-КОНСУЛЬТАНТОМ ПО ДАМПУ*\n\n"
            f"• *Г¤ш¤ф◊ ¶‡ ∆◊§∂¤:* `{proj['name']}`\n"
            f"• *П§дш†◊§ бум‡г †з д‡мп‡:* *{len(proj['files'])} ш◊.*\n"
            f"• *К∂юч¤й ∑ §б§йм¤:* *{len(keys)} ш◊.*\n\n"
            f"✍️ *Сп∏‡ш†∑‡й◊¤ ш§ уг§д¶§ †∂† к†д‡й◊¤ ∆к∏†¶ш§◊* — к§¶∆у∂ь◊‡¶◊ уж¤ д¤∏ж†◊ п¤∏¤д г∂‡з‡м† ∑∆¤ ∑‡ш† ф‡й∂ы † к§ды!"
        ),
        reply_markup=kb
    )

@dp.callback_query(F.data.startswith("fai_"))
async def _cb_file_ai(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    await callback.answer()
    await _view_ai_session(callback.message, u, state)

@dp.callback_query(F.data == "ai_clear_hist", StateFilter("*"))
async def _cb_ai_clear(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    proj["ai_history"] = []
    _db_write()
    await callback.answer(_u("П∏§◊§к§∂ б¤∆¤ды ч†∆◊, к‡к ∆∂¤з‡!"), show_alert=True)

@dp.message(SessionCtx.s_aichat, F.text)
async def _step_ai_msg(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    locked, lock_msg = _check_gate(u)
    if locked:
        await _send(message, lock_msg)
        return
    proj = _ctx_proj(u)
    user_q = message.text.strip()

    wait_msg = await _send(message, _u(f"⏳ Т‡к† ∂†∆◊‡¤м ∑‡ш д‡мп «{proj['name']}» ({len(proj['files'])} ф‡й∂.) † г§◊§∑†м §◊∑¤◊..."))
    answer_text = await _query_llm_pool(u, proj, user_q)
    if wait_msg:
        try:
            await wait_msg.delete()
        except Exception:
            pass

    proj.setdefault("ai_history", []).append({"role": "user", "text": user_q})
    proj["ai_history"].append({"role": "model", "text": answer_text})
    _db_write()

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=_u("💾 П§дш†◊ь э◊§◊ §◊∑¤◊ ∑ Д‡мп"), callback_data="ai_save_last")]
    ])
    await _send_chunked(message, answer_text)
    await _send(message, _u("👆 Ж¤∂‡¤◊¤ ∆§х∏‡¶†◊ь §◊∑¤◊ д∂я †∆◊§∏††?"), reply_markup=kb)

@dp.callback_query(F.data == "ai_save_last", StateFilter("*"))
async def _cb_ai_save(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    proj = _ctx_proj(u)
    if proj.get("ai_history"):
        last_ans = proj["ai_history"][-1]["text"]
        title = _u(f"К§¶∆у∂ь◊‡ц†я ИИ ({datetime.now().strftime('%H:%M')})")
        proj["files"].append({
            "type": "note",
            "file_id": None,
            "name": title,
            "caption": "",
            "extracted_text": last_ans,
            "pages": _paginate_str(last_ans),
            "date": datetime.now().strftime("%d.%m %H:%M")
        })
        _db_write()
        await callback.answer(_u("✅ П§дш†◊§ ∑ д‡мп п∏§¤к◊‡!"), show_alert=True)

async def _view_keys(message: types.Message, u: dict):
    p_keys = u.get("personal_keys", [])
    btns = []
    for idx, k in enumerate(p_keys):
        masked = f"{k[:8]}...{k[-4:]}" if len(k) > 14 else k
        btns.append([InlineKeyboardButton(text=_u(f"🔑 К∂юч #{idx+1} ({masked}) — 🗑 Уб∏‡◊ь"), callback_data=f"kdel_{idx}")])

    btns.append([InlineKeyboardButton(text=_u("➕ П§дш†◊ь ¶§∑ый к∂юч(†) API"), callback_data="kadd_p")])
    if p_keys:
        btns.append([InlineKeyboardButton(text=_u("🧹 Оч†∆◊†◊ь ∑∆ю ∆∑язку к∂юч¤й"), callback_data="kclear_all")])

    await _send(
        message,
        _u(
            f"`{u['ui']['status_badge']}`\n"
            f"🔑 *ВАША ЛИЧНАЯ СВЯЗКА КЛЮЧЕЙ API*\n\n"
            f"• *К∂юч¤й ∑ ∑‡ш¤й §б§йм¤:* *{len(p_keys)} ш◊.*\n\n"
            f"💡 _З‡к†ды∑‡й◊¤ ∆юд‡ §◊ 1 д§ 4 к∂юч¤й (п§дд¤∏ж†∑‡¤◊∆я ∂юб§й ф‡∆§¶: `AQ.Ab8...`, `AIza...` † д∏.). "
            f"К§гд‡ §д†¶ к∂юч у∆◊‡¶¤◊, б§◊ ∆‡м п¤∏¤к∂юч†◊∆я ¶‡ ∆∂¤дующ†й б¤з шум‡ † пы∂†!_"
        ),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=btns)
    )

@dp.callback_query(F.data == "kadd_p")
async def _cb_key_add(callback: types.CallbackQuery, state: FSMContext):
    await state.set_state(SessionCtx.s_akey)
    await _send(callback.message, _u("🔑 О◊п∏‡∑ь◊¤ §д†¶ †∂† ∆∏‡зу п‡чку к∂юч¤й (к‡ждый ∆ ¶§∑§й ∆◊∏§чк†):"), reply_markup=_abort_kb())
    await callback.answer()

@dp.message(SessionCtx.s_akey)
async def _step_key_add(message: types.Message, state: FSMContext):
    await state.clear()
    u = _ctx_user(message.from_user)
    raw_lines = [line.strip() for line in (message.text or "").splitlines() if len(line.strip()) >= 10]
    if not raw_lines:
        await _send(message, _u("⚠️ Ш§-◊§ э◊§◊ к∂юч ∆§∑∆¤м к§∏§◊к†й, п∏§∑¤∏ь◊¤ ¤щё ∏‡з."), reply_markup=_build_reply_kb(u))
        return
    target = u.setdefault("personal_keys", [])
    added = 0
    for k in raw_lines:
        if k not in target:
            target.append(k)
            added += 1
    _db_write()
    await _send(message, _u(f"✅ Д§б‡∑∂¤¶§ к∂юч¤й ∑ §б§йму: *{added} ш◊.*!"), reply_markup=_build_reply_kb(u))
    await _view_keys(message, u)

@dp.callback_query(F.data.startswith("kdel_"))
async def _cb_key_del(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    idx = int(callback.data.replace("kdel_", ""))
    if 0 <= idx < len(u.get("personal_keys", [])):
        u["personal_keys"].pop(idx)
        _db_write()
        await callback.answer(_u("К∂юч уб∏‡¶"))
        await _view_keys(callback.message, u)

@dp.callback_query(F.data == "kclear_all")
async def _cb_key_clear(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    u["personal_keys"] = []
    u["key_cursor"] = 0
    _db_write()
    await callback.answer(_u("С∑язк‡ к∂юч¤й §ч†щ¤¶‡"))
    await _view_keys(callback.message, u)

@dp.message(SessionCtx.s_search)
async def _step_search(message: types.Message, state: FSMContext):
    await state.clear()
    u = _ctx_user(message.from_user)
    q = (message.text or "").strip().lower()
    if not q:
        return
    found = []
    for _, p in u["projects"].items():
        for fitem in p["files"]:
            txt = fitem.get("extracted_text", "")
            if q in fitem["name"].lower() or q in txt.lower():
                pos = txt.lower().find(q)
                snippet = txt[max(0, pos - 60):pos + 140].replace("\n", " ") if pos != -1 else fitem.get("caption", "")
                found.append(f"📁 *{p['name']}* $\\rightarrow$ 📄 *{fitem['name']}*:\n`...{snippet}...`")
    for c in u["my_codes"] + db["public_codes"]:
        if q in c["title"].lower() or q in c.get("code", "").lower() or q in c.get("desc", "").lower():
            found.append(_u(f"💻 *К§д: {c['title']}*\n_{c.get('desc', '')}_"))
    if not found:
        await _send(message, _u(f"📭 И∆к‡∂†-†∆к‡∂† `{q}`, ¶§ ◊‡к† ¶†ч¤г§ ¶¤ ¶‡ш∂†."), reply_markup=_build_reply_kb(u))
        return
    await _send(
        message,
        _u(f"🔍 *См§◊∏†◊¤, ш§ мы ¶‡ш∂† ∑ з‡к∏§м‡х ({len(found)}):*\n\n") + "\n\n---\n\n".join(found[:10]),
        reply_markup=_build_reply_kb(u)
    )

async def _view_settings(message: types.Message, u: dict, page: int = 0, edit: bool = False):
    ui = u["ui"]
    items = list(UI_LABELS.items())
    per_page = 4
    total_pages = max(1, math.ceil(len(items) / per_page))
    page = max(0, min(page, total_pages - 1))
    subset = items[page * per_page : (page + 1) * per_page]

    btns = []
    for k, label in subset:
        curr_val = ui.get(k, "")[:24]
        btns.append([InlineKeyboardButton(text=f"✏️ {label}: «{curr_val}»", callback_data=f"uied_{k}")])

    nav = []
    if page > 0:
        nav.append(InlineKeyboardButton(text="⬅️", callback_data=f"spage_{page-1}"))
    nav.append(InlineKeyboardButton(text=_u(f"К¶§пк† {page+1}/{total_pages}"), callback_data="r_noop"))
    if page < total_pages - 1:
        nav.append(InlineKeyboardButton(text="➡️", callback_data=f"spage_{page+1}"))
    btns.append(nav)

    curr_ps = u.get("page_size", 5)
    btns.append([
        InlineKeyboardButton(text=_u(f"📄 П‡г†¶‡ц†я ∆п†∆к§∑: п§ {curr_ps} ш◊."), callback_data="s_toggle_psize"),
        InlineKeyboardButton(text=_u("🔄 Сб∏§∆ ¶‡з∑‡¶†й к¶§п§к"), callback_data="ui_reset")
    ])
    btns.append([
        InlineKeyboardButton(text=_u("ℹ️ О †∆п§∂ьз§∑‡¶†† (Сб∏§∆ / П‡пк†)"), callback_data="open_usage_sec")
    ])

    txt = _u(
        f"`{ui['status_badge']}`\n"
        f"⚙️ *НАСТРОЙКИ КНОПОК И СПИСКОВ*\n\n"
        f"• П¤∏¤†м¤¶§∑ы∑‡й◊¤ ∂юбы¤ к¶§пк† † ∆◊∏§ку `{ui['status_badge']}` п§д ∆∑§й ∑ку∆.\n"
        f"• Н‡∆◊∏‡†∑‡й◊¤ п‡г†¶‡ц†ю ∆п†∆к§∑ (∆к§∂ьк§ ф‡й∂§∑ † к§д§∑ п§к‡зы∑‡◊ь ¶‡ §д¶§й ∆◊∏‡¶†ц¤).\n"
        f"• П¤∏¤х§д м¤жду п‡пк‡м† † к¶§пк† ∆б∏§∆‡ ¶‡х§дя◊∆я ∑ ∏‡зд¤∂¤ *«ℹ️ О †∆п§∂ьз§∑‡¶††»*:"
    )
    kb = InlineKeyboardMarkup(inline_keyboard=btns)
    if edit:
        await _edit(message, txt, reply_markup=kb)
    else:
        await _send(message, txt, reply_markup=kb)

@dp.callback_query(F.data.startswith("spage_"))
async def _cb_settings_page(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    page = int(callback.data.replace("spage_", ""))
    await _view_settings(callback.message, u, page=page, edit=True)
    await callback.answer()

@dp.callback_query(F.data == "s_toggle_psize")
async def _cb_toggle_psize(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    cycle = {5: 8, 8: 12, 12: 5}
    u["page_size"] = cycle.get(u.get("page_size", 5), 5)
    _db_write()
    await callback.answer(_u(f"П‡г†¶‡ц†я ∆п†∆к§∑: п§ {u['page_size']} э∂¤м¤¶◊§∑ ¶‡ ∆◊∏‡¶†ц¤"))
    await _view_settings(callback.message, u, page=0, edit=True)

@dp.callback_query(F.data.startswith("uied_"))
async def _cb_ui_edit(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    ui_key = callback.data.replace("uied_", "")
    if ui_key in UI_LABELS:
        await state.update_data(ui_key=ui_key)
        await state.set_state(SessionCtx.s_uival)
        await _send(
            callback.message,
            _u(f"✏️ Т¤кущ¤¤ з¶‡ч¤¶†¤ ({UI_LABELS[ui_key]}):\n`{u['ui'].get(ui_key, '')}`\n\nО◊п∏‡∑ь◊¤ ∑‡ш ¶§∑ый ∑‡∏†‡¶◊:"),
            reply_markup=_abort_kb()
        )
    await callback.answer()

@dp.message(SessionCtx.s_uival)
async def _step_ui_edit(message: types.Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    u = _ctx_user(message.from_user)
    ui_key = data.get("ui_key")
    if ui_key and message.text:
        u["ui"][ui_key] = message.text.strip()
        _db_write()
    await _send(message, _u("✅ Вы∑¤∆к‡ §б¶§∑∂¤¶‡ п§д ∑‡ш ∑ку∆!"), reply_markup=_build_reply_kb(u))
    await _view_settings(message, u, page=0, edit=False)

@dp.callback_query(F.data == "ui_reset")
async def _cb_ui_reset(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    u["ui"] = dict(DEFAULT_UI)
    _db_write()
    await callback.answer(_u("З‡∑§д∆к†¤ ¶‡з∑‡¶†я к¶§п§к ∑§з∑∏‡щ¤¶ы"))
    await _send(callback.message, _u(f"🔄 С◊‡¶д‡∏◊¶ы¤ ¶‡з∑‡¶†я к¶§п§к {APP_VERSION} ∑§∆∆◊‡¶§∑∂¤¶ы!"), reply_markup=_build_reply_kb(u))
    await _view_settings(callback.message, u, page=0, edit=False)

async def _view_usage_info(message: types.Message, u: dict, edit: bool = False):
    grp = u.get("group", "n0")
    if grp == "n1":
        txt = _u(
            f"🌲 *О ИСПОЛЬЗОВАНИИ • ПАМЯТКА ПУТНИКУ*\n"
            f"📦 *В¤∏∆†я ∆б§∏к†:* `{APP_VERSION}` • *Д‡◊‡ ∆б§∏к†:* `{BUILD_DATE}`\n\n"
            f"*З‡ч¤м э◊§ м¤∆◊§:*\n"
            f"* (Вы ¶‡х§д†◊¤∆ь ¶‡ ◊†х§й з‡†мк¤ ∆∏¤д† у∏‡∂ь∆к†х ∆§∆¤¶ † ∆◊‡∏ых ∏‡д†§∑ыш¤к. "
            f"Зд¤∆ь м§ж¶§ ∆∂уш‡◊ь ∆∑§дк† †з-з‡ х∏¤б◊‡, п∏§пу∆к‡◊ь ∆∂§∑‡ ∆к∑§зь ∂‡мп§∑ый д¤к§д¤∏ эф†∏‡ "
            f"†∂† ∑¤∏¶у◊ь∆я §б∏‡◊¶§ к ∆∑§†м ∏‡б§ч†м п‡пк‡м.)\n\n"
            f"*З¶‡ч¤¶†я к¶§п§к ∑¶†зу:*\n"
            f"• *«🛠 О◊ з‡∑†∆‡¶†й»* — ∆◊∏ях†∑‡¤◊ п§м¤х† эф†∏‡, ∆б∏‡∆ы∑‡¤◊ з‡∆◊∏я∑ш†¤ д†‡∂§г† † ∑§з∑∏‡щ‡¤◊ ∆∑§б§ду ш‡г‡, ∆§х∏‡¶яя ∑∆¤ ∑‡ш† ¶‡х§дк†.\n"
            f"• *«🧨 Сб∏§∆»* — п§∂¶§∆◊ью ∆◊†∏‡¤◊ ∑∆¤ ∑‡ш† ∆∂¤ды, ф‡й∂ы, к§ды † з‡п†∆† †з б‡зы д‡¶¶ых ∑ ‡б∆§∂ю◊¶ый ¶§∂ь, буд◊§ ∑ы ¶†к§гд‡ ∆юд‡ ¶¤ з‡х§д†∂†."
        )
        rows = [
            [
                InlineKeyboardButton(text=_u("🛠 О◊ з‡∑†∆‡¶†й"), callback_data="act_unstick"),
                InlineKeyboardButton(text=_u("🧨 Сб∏§∆"), callback_data="act_full_wipe")
            ],
            [
                InlineKeyboardButton(text=_u("📁 П‡пк‡: П∏§¤к◊ы"), callback_data="nav_dir_0"),
                InlineKeyboardButton(text=_u("🌲 П‡пк‡: Л¤∆ (Ак◊†∑¶‡)"), callback_data="nav_dir_1")
            ]
        ]
    else:
        txt = _u(
            f"`{u['ui']['status_badge']}`\n"
            f"ℹ️ *НАСТРОЙКИ / О ИСПОЛЬЗОВАНИИ*\n"
            f"📦 *В¤∏∆†я ∆б§∏к†:* `{APP_VERSION}` • *Д‡◊‡ §б¶§∑∂¤¶†я:* `{BUILD_DATE}`\n\n"
            f"*З‡ч¤м ¶уж¤¶ э◊§◊ б§◊:*\n"
            f"Т‡к† ∆∂уш‡й◊¤: э◊§ ∑‡ш ∂†ч¶ый ‡∑◊§¶§м¶ый Г¤ш¤ф◊-Х‡б! Сюд‡ ∑ы ∆к∂‡ды∑‡¤◊¤ д‡мпы ∂юбых д§кум¤¶◊§∑ "
            f"(м¤◊§д†чк†, п∏¤з¤¶◊‡ц††, ‡∏х†∑ы † ∆к∏†¶ы), ч†◊‡¤◊¤ †х п∏ям§ п§ ¶‡ж‡◊†ю ¶‡ ф‡й∂, д¤∏ж†◊¤ ∆∑§† к§ды "
            f"∑ ∂†ч¶§м ∆¤йф¤ †∂† ¶‡ §бщ¤м П∏†∑§з¤ † §бщ‡¤◊¤∆ь ∆ ИИ-к§¶∆у∂ь◊‡¶◊§м ∆∏‡зу п§ ∑∆¤му д‡мпу ч¤∏¤з ∆∑язку ∆∑§†х API-к∂юч¤й.\n\n"
            f"*Оп†∆‡¶†¤ к¶§п§к ∆¶†зу:*\n"
            f"• *«🛠 О◊ з‡∑†∆‡¶†й»* — ∆к§∏‡я п§м§щь §◊ ∂юбых ∆§ф◊∂§к§∑: ∆б∏‡∆ы∑‡¤◊ з‡∑†∆ш†й ∑∑§д, ◊‡йм¤∏ы † ∆§∆◊§я¶†я, ¶¤ ◊∏§г‡я ∑‡ш† ф‡й∂ы † к§ды.\n"
            f"• *«🧨 Сб∏§∆»* — п§∂¶‡я ∂†к∑†д‡ц†я ‡кк‡у¶◊‡ ∑ ¶§∂ь: уд‡∂я¤◊ ∑‡∆ †з б‡зы д‡¶¶ых † ∆◊‡◊†∆◊†к†, ∆◊†∏‡¤◊ ∑∆¤ ∑‡ш† п∏§¤к◊ы, д‡мпы † к§ды, буд◊§ ∑‡∆ ◊у◊ ¶†к§гд‡ † ¶¤ ∆◊§я∂§."
        )
        rows = [
            [
                InlineKeyboardButton(text=_u("🛠 О◊ з‡∑†∆‡¶†й"), callback_data="act_unstick"),
                InlineKeyboardButton(text=_u("🧨 Сб∏§∆"), callback_data="act_full_wipe")
            ],
            [
                InlineKeyboardButton(text=_u("📁 П‡пк‡: П∏§¤к◊ы (Ак◊†∑¶‡)"), callback_data="nav_dir_0"),
                InlineKeyboardButton(text=_u("🌲 П‡пк‡: Л¤∆"), callback_data="nav_dir_1")
            ],
            [InlineKeyboardButton(text=_u("⬅️ Н‡з‡д ∑ ¶‡∆◊∏§йк†"), callback_data="back_to_settings")]
        ]

    kb = InlineKeyboardMarkup(inline_keyboard=rows)
    if edit:
        await _edit(message, txt, reply_markup=kb)
    else:
        await _send(message, txt, reply_markup=kb)

@dp.callback_query(F.data == "open_usage_sec")
async def _cb_open_usage(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    await _view_usage_info(callback.message, u, edit=True)
    await callback.answer()

@dp.callback_query(F.data == "back_to_settings")
async def _cb_back_settings(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    await _view_settings(callback.message, u, page=0, edit=True)
    await callback.answer()

@dp.callback_query(F.data == "act_unstick", StateFilter("*"))
async def _cb_unstick(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    _stop_worker(u["id"])
    await state.clear()
    u["rt"] = {}
    if u.get("group") == "n2":
        u["group"] = "n0"
    _db_write()
    await callback.answer(_u("✅ В∆¤ з‡∑†∆‡¶†я † ∆§∆◊§я¶†я ∆б∏§ш¤¶ы!"), show_alert=True)
    await _send(
        callback.message,
        _u("🛠 *З‡∑†∆‡¶†я ∆¶я◊ы!* В∆¤ ∑‡ш† ф‡й∂ы † к§ды ∑ п§∂¶§й ∆§х∏‡¶¶§∆◊†, ∏‡б§◊‡¤м д‡∂ьш¤:"),
        reply_markup=_build_reply_kb(u)
    )

@dp.callback_query(F.data == "act_full_wipe", StateFilter("*"))
async def _cb_full_wipe(callback: types.CallbackQuery, state: FSMContext):
    await state.clear()
    _wipe_uid(str(callback.from_user.id))
    await callback.answer(_u("П∏§ф†∂ь п§∂¶§∆◊ью ∆◊ё∏◊ ∑ 0."), show_alert=True)
    await _edit(
        callback.message,
        _u(
            "🧹 *П§∂¶ый ∆б∏§∆ ∑ып§∂¶¤¶.*\n"
            "В∆¤ ∑‡ш† д‡¶¶ы¤, д‡мпы, к§ды † ∆◊‡◊†∆◊†к‡ п§∂¶§∆◊ью уд‡∂¤¶ы †з б‡зы д‡¶¶ых ∑ 0.\n"
            "Ч◊§бы ¶‡ч‡◊ь ∆ ч†∆◊§г§ ∂†∆◊‡ — ¶‡жм†◊¤ /start."
        )
    )

@dp.callback_query(F.data.in_(["nav_dir_0", "nav_dir_1"]), StateFilter("*"))
async def _cb_switch_dir(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    target = callback.data.replace("nav_dir_", "")
    _stop_worker(u["id"])
    await state.clear()

    if target == "0":
        u["group"] = "n0"
        if _u("Л¤∆") in u["ui"]["status_badge"]:
            u["ui"]["status_badge"] = _u(f"🧭 Г∏упп‡: П∏§¤к◊ы • Сб§∏к‡ {APP_VERSION}")
        _db_write()
        await callback.answer(_u("П¤∏¤х§д ∑ п‡пку: П∏§¤к◊ы"))
        await _cmd_entry(callback.message, state)
    else:
        u["group"] = "n1"
        if _u("П∏§¤к◊ы") in u["ui"]["status_badge"]:
            u["ui"]["status_badge"] = _u(f"🧭 Г∏упп‡: Л¤∆ • Сб§∏к‡ {APP_VERSION}")
        _db_write()
        await callback.answer(_u("П¤∏¤х§д ∑ п‡пку: Л¤∆"))
        await _view_n1(callback.message, u, "e")

def _mutate_token(seed_w: str) -> str:
    g_pool = _R["g"]
    leet = {"e": "3", "a": "4", "o": "0", "t": "7", "\u0430": "4", "\u043e": "0", "\u0435": "3", "\u0440": "p", "\u043d": "H", "\u0443": "y"}
    ch_list = list(_u(seed_w))
    for i in range(len(ch_list)):
        if random.random() < 0.28 and ch_list[i].lower() in leet:
            ch_list[i] = leet[ch_list[i].lower()]
    pos = random.randint(2, max(2, len(ch_list) - 2))
    ch_list.insert(pos, random.choice(g_pool))
    if random.random() < 0.65:
        pos2 = random.randint(1, len(ch_list) - 1)
        ch_list.insert(pos2, random.choice(["_", "/", "~", "・", "4", "7"]))
    return "".join(ch_list)[:28]

async def _view_n1(message: types.Message, u: dict, mode: str, state: FSMContext = None):
    pmap = {"e": _R["b0"], "d": _R["b1"], "n": _R["b2"]}
    chosen = _u(random.choice(pmap.get(mode, _R["b0"])))
    s_btn = _mutate_token(random.choice(_R["ks"]))

    rows = []
    if mode == "e":
        rows.append([
            InlineKeyboardButton(text=_u("* О∆м§◊∏¤◊ь м‡∂‡х†◊§∑ый м§х"), callback_data="n1_act_e"),
            InlineKeyboardButton(text=_u("* П§∆∂уш‡◊ь гу∂ узк§к§∂¤йк†"), callback_data="n1_act_n")
        ])
    elif mode == "d":
        if state:
            await state.set_state(SessionCtx.s_n1_dec)
        rows.append([InlineKeyboardButton(text=_u("📻 Н‡∆◊∏§†◊ь ч‡∆◊§◊у ¤щё ∏‡з"), callback_data="n1_act_d")])
    elif mode == "n":
        rows.append([InlineKeyboardButton(text=_u("📜 П¤∏¤∑¤∏¶у◊ь б¤∏¤∆◊я¶§й ∂†∆◊"), callback_data="n1_act_n")])

    pair = [
        InlineKeyboardButton(text=_u(_R["kn"]), callback_data="n1_norm"),
        InlineKeyboardButton(text=s_btn, callback_data="n2_boot")
    ]
    random.shuffle(pair)
    rows.append(pair)

    await _send(message, chosen, reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))

@dp.callback_query(F.data.startswith("n1_act_"))
async def _cb_n1_act(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    mode = callback.data.replace("n1_act_", "")
    await callback.answer()
    await _view_n1(callback.message, u, mode, state)

@dp.callback_query(F.data == "n1_norm")
async def _cb_n1_norm(callback: types.CallbackQuery):
    await callback.answer(_u("* (К‡ж¤◊∆я, ∑∆ё ∆п§к§й¶§. Тум‡¶ ∆◊§†◊ ¶¤п§д∑†ж¶§.)"), show_alert=False)

@dp.message(SessionCtx.s_n1_dec, F.text)
async def _step_n1_echo(message: types.Message, state: FSMContext):
    u = _ctx_user(message.from_user)
    if _check_gate(u)[0]:
        return
    txt = message.text.strip()
    glitched = "".join(c + random.choice(["", "", "░", "·", "~"]) for c in txt[::-1])
    await _send(
        message,
        _u(
            f"* (Д¤к§д¤∏ х∏†п†◊ ◊ёп∂ым ∂‡мп§∑ым эх§м †з-п§д ◊§∏фя¶†к‡:)\n`{glitched}`\n\n"
            f"* (Гд¤-◊§ ∑ г∂уб†¶¤ ш‡х◊ы эх§-ц∑¤◊§к п§∑◊§∏†∂ ∑‡ш† ∆∂§∑‡ † з‡◊†х.)"
        ),
        reply_markup=_build_reply_kb(u)
    )

def _freeze_30m(u: dict):
    _stop_worker(u["id"])
    u["lock_until"] = time.time() + 1800
    u["lock_reason"] = "30m"
    u["group"] = "n0"
    u["rt"] = {}
    _db_write()

async def _init_n2(message: types.Message, u: dict):
    _stop_worker(u["id"])
    u["group"] = "n2"
    u["rt"] = {
        "p": "p1",
        "tr": random.randint(5, 9),
        "cr": 0
    }
    _db_write()
    await _render_p1(message, u, edit=False)

async def _render_p1(message: types.Message, u: dict, edit: bool = True):
    st = u["rt"]
    mono = _u(random.choice(_R["ci"]))
    t_lbl = _mutate_token(random.choice(_R["ks"]))

    c_pos = random.randint(0, 5)
    kb_rows, row = [], []
    for i in range(6):
        if i == c_pos:
            row.append(InlineKeyboardButton(text=t_lbl, callback_data="n2_p1_ok"))
        else:
            row.append(InlineKeyboardButton(text=_u(_R["kn"]), callback_data=f"n2_p1_er_{i}"))
        if len(row) == 2:
            kb_rows.append(row)
            row = []

    text = f"...\n\n_{mono}_\n\n`[{st['cr'] + 1}]`"
    kb = InlineKeyboardMarkup(inline_keyboard=kb_rows)
    if edit:
        await _edit(message, text, reply_markup=kb)
    else:
        await _send(message, text, reply_markup=kb)

@dp.callback_query(F.data == "n2_boot", StateFilter("*"))
async def _cb_n2_boot(callback: types.CallbackQuery, state: FSMContext):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    await state.clear()
    await callback.answer()
    await _init_n2(callback.message, u)

@dp.callback_query(F.data.startswith("n2_p1_er_"), StateFilter("*"))
async def _cb_n2_p1_er(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    _freeze_30m(u)
    await callback.answer()
    await _edit(
        callback.message,
        _u("...◊ы ∑ыб∏‡∂ п∏†◊∑§∏†◊ь∆я, ч◊§ ∑∆ё ¶§∏м‡∂ь¶§. ¶§ ∑¶у◊∏† б§∂ьш¤ ¶†ч¤г§ ¶¤◊.\n\n_(д§∆◊уп з‡б∂§к†∏§∑‡¶ ¶‡ 30 м†¶у◊)_")
    )

@dp.callback_query(F.data == "n2_p1_ok", StateFilter("*"))
async def _cb_n2_p1_ok(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    st = u.get("rt", {})
    if st.get("p") != "p1":
        return

    st["cr"] = st.get("cr", 0) + 1
    if st["cr"] >= st.get("tr", 5):
        st["p"] = "p2"
        st["hp"] = random.randint(2, 10)
        st["a"] = 0
        st["ga"] = random.choice([i for i in range(1, 9)])
        st["gc"] = random.choice(_R["g"][:15])
        st["gs"] = False
        st["fr"] = "0"
        st["bt"] = 0
        _db_write()
        await callback.answer()
        await _render_p2(callback.message, u)
        _spawn_p2_ticker(callback.message, u)
    else:
        _db_write()
        await callback.answer()
        await _render_p1(callback.message, u, edit=True)

def _vx_mesh(a_idx: int, beat: int, g_ch: str | None, fr_mode: str) -> str:
    theta = math.radians(a_idx * 40)
    sin_t, cos_t = math.sin(theta), math.cos(theta)
    scale = 1.06 if beat == 1 else 0.96
    shades = "@#%&*+=-:." if beat == 1 else "#%&*+=-:.."

    w, h = 27, 13
    lines = []
    for r in range(h):
        y = (6.0 - r) / 4.2 / scale
        row_c = []
        for c in range(w):
            x = (c - 13.0) / 7.2 / scale
            rx = x * cos_t
            in_top = (0.75 < y <= 1.38) and (abs(x - 0.15 * cos_t) < 0.22 or abs(x + 0.32 * cos_t) < 0.16)
            hx = x - 0.08 * y * cos_t
            eq = (hx * hx + y * y - 0.85) ** 3 - (hx * hx) * (y ** 3) * 1.35

            if in_top:
                row_c.append("|" if beat == 0 else "H")
            elif eq <= 0.0:
                depth = max(0.0, min(1.0, math.sqrt(max(0.0, 1.0 - min(1.0, hx * hx + y * y)))))
                groove = math.sin((x * 3.2 - y * 1.8) + theta)
                lum = 0.55 * depth + 0.25 * (rx * sin_t + depth * cos_t) + 0.20 * groove
                s_idx = max(0, min(len(shades) - 1, int((1.0 - max(0.0, min(1.0, lum))) * (len(shades) - 1))))
                row_c.append(shades[s_idx])
            else:
                row_c.append(" ")
        lines.append("".join(row_c))

    if g_ch:
        mid_r = 6
        mid_c = 13 + int(2 * sin_t)
        r_list = list(lines[mid_r])
        r_list[mid_c - 1:mid_c + 2] = ["[", g_ch, "]"]
        lines[mid_r] = "".join(r_list)

    if fr_mode == "r":
        b = "\U0001f534" * 15
        return "\n".join([b] + [f"\U0001f534`{ln}`\U0001f534" for ln in lines] + [b])
    if fr_mode == "g":
        b = "\u2747\ufe0f" * 15
        return "\n".join([b] + [f"\u2747\ufe0f`{ln}`\u2747\ufe0f" for ln in lines] + [b])
    return "```text\n" + "\n".join(lines) + "\n```"

def _p2_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="⬅️ -40°", callback_data="n2_rl"),
            InlineKeyboardButton(text="➡️ +40°", callback_data="n2_rr")
        ],
        [
            InlineKeyboardButton(text="\U0001f5e1", callback_data="n2_ca"),
            InlineKeyboardButton(text="\U0001f6e1", callback_data="n2_cd"),
            InlineKeyboardButton(text="\u2733\ufe0f", callback_data="n2_cp")
        ]
    ])

async def _render_p2(message: types.Message, u: dict):
    st = u.get("rt", {})
    if st.get("p") != "p2":
        return
    a_idx = st.get("a", 0)
    deg = a_idx * 40
    beat = st.get("bt", 0)
    fr = st.get("fr", "0")
    g_ch = st["gc"] if a_idx == st.get("ga") else None

    mesh = _vx_mesh(a_idx, beat, g_ch, fr)
    ico = "\U0001fac0" if beat == 1 else "\U0001f5a4"
    wh = _u(_R["cw"][a_idx % len(_R["cw"])])

    await _edit(message, f"{ico} `{deg}° / 360°`\n_{wh}_\n\n{mesh}", reply_markup=_p2_kb())

def _spawn_p2_ticker(message: types.Message, u: dict):
    uid = u["id"]
    _stop_worker(uid)

    async def _loop():
        try:
            while True:
                await asyncio.sleep(2.3)
                if _check_gate(u)[0]:
                    break
                st = u.get("rt", {})
                if st.get("p") != "p2":
                    break
                st["bt"] = 1 - st.get("bt", 0)
                await _render_p2(message, u)
        except asyncio.CancelledError:
            pass

    _TASKS[uid] = asyncio.create_task(_loop())

async def _eject_to_n0(message: types.Message, u: dict):
    _stop_worker(u["id"])
    u["group"] = "n0"
    u["rt"] = {}
    _db_write()
    await _edit(
        message,
        _u("...◊ы ¶¤ ∆◊‡∂ м§∂†◊ь∆я, к§гд‡ з¤∂ё¶ый ∆∑¤◊ \u2747\ufe0f ∆§мк¶у∂∆я ∑§к∏уг пу∆◊§◊ы. ◊¤бя ∑ыб∏§∆†∂§ §б∏‡◊¶§ ¶‡ п§∑¤∏х¶§∆◊ь.")
    )
    await _send(message, _u(f"`{u['ui']['status_badge']}`\nВы ∑§з∑∏‡щ¤¶ы ∑ ∏‡б§чую п‡пку."), reply_markup=_build_reply_kb(u))

@dp.callback_query(F.data.in_(["n2_rl", "n2_rr"]), StateFilter("*"))
async def _cb_k_rot(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    st = u.get("rt", {})
    if st.get("p") != "p2":
        return

    if st.get("fr") == "r":
        _freeze_30m(u)
        await callback.answer()
        await _edit(callback.message, _u("...б‡г∏§∑‡я ∏‡мк‡ \U0001f534 ∆ж‡∂‡∆ь. ◊ы ¶¤ з‡щ†◊†∂∆я. б∂§к†∏§∑к‡ ¶‡ 30 м†¶у◊."))
        return
    if st.get("fr") == "g":
        await callback.answer()
        await _eject_to_n0(callback.message, u)
        return
    if st.get("gs"):
        _freeze_30m(u)
        await callback.answer()
        await _edit(callback.message, _u("...з¶‡к м¤∂ьк¶у∂ ∂†шь ∏‡з, † ◊ы §◊∑¤∏¶у∂∆я, ¶¤ ¶‡¶¤∆я уд‡∏. б∂§к†∏§∑к‡ ¶‡ 30 м†¶у◊."))
        return

    step = -1 if callback.data == "n2_rl" else 1
    st["a"] = (st.get("a", 0) + step) % 9
    st["bt"] = 1 - st.get("bt", 0)

    if st["a"] == st.get("ga"):
        st["gs"] = True
        st["fr"] = "0"
    else:
        roll = random.random()
        st["fr"] = "r" if roll < 0.27 else ("g" if roll < 0.48 else "0")

    _db_write()
    await callback.answer()
    await _render_p2(callback.message, u)

@dp.callback_query(F.data.in_(["n2_ca", "n2_cd", "n2_cp"]), StateFilter("*"))
async def _cb_k_act(callback: types.CallbackQuery):
    u = _ctx_user(callback.from_user)
    if _check_gate(u)[0]:
        return
    st = u.get("rt", {})
    if st.get("p") != "p2":
        return

    act = callback.data.replace("n2_c", "")
    fr = st.get("fr", "0")
    on_g = (st.get("a") == st.get("ga"))

    if act == "a":
        if fr == "r":
            _freeze_30m(u)
            await callback.answer()
            await _edit(callback.message, _u("...◊ы з‡м‡х¶у∂∆я ∑§ ∑∏¤мя †мпу∂ь∆‡ \U0001f534 ∑м¤∆◊§ з‡щ†◊ы. б∂§к†∏§∑к‡ ¶‡ 30 м†¶у◊."))
            return
        if fr == "g":
            await callback.answer()
            await _eject_to_n0(callback.message, u)
            return
        if not on_g:
            _freeze_30m(u)
            await callback.answer()
            await _edit(callback.message, _u("...к∂†¶§к уд‡∏†∂ ∑ пу∆◊§◊у б¤з з¶‡к‡. §д¶‡ §ш†бк‡ — 30 м†¶у◊ ◊†ш†¶ы."))
            return

        st["hp"] -= 1
        st["bt"] = 1 - st.get("bt", 0)
        if st["hp"] <= 0:
            _stop_worker(u["id"])
            st["p"] = "p3"
            _db_write()
            await callback.answer(_u("...§¶§ ∏‡∆к§∂§∂§∆ь."))
            await _run_p3_seq(callback.message, u)
            return

        curr_a = st["a"]
        st["ga"] = random.choice([i for i in range(9) if i != curr_a])
        st["gc"] = random.choice(_R["g"][:15])
        st["gs"] = False
        st["fr"] = "0"
        _db_write()
        await callback.answer(_u("...уд‡∏ п∏§шё∂ ∆к∑§зь з¶‡к."))
        await _render_p2(callback.message, u)
        return

    if act == "d":
        if fr == "r":
            st["fr"] = "0"
            st["bt"] = 1 - st.get("bt", 0)
            _db_write()
            await callback.answer(_u("...†мпу∂ь∆ \U0001f534 §◊∏‡жё¶."))
            await _render_p2(callback.message, u)
            return
        if fr == "g":
            await callback.answer()
            await _eject_to_n0(callback.message, u)
            return
        if st.get("gs"):
            _freeze_30m(u)
            await callback.answer()
            await _edit(callback.message, _u("...з¶‡к жд‡∂ уд‡∏‡, ¶§ ◊ы з‡к∏ы∂∆я щ†◊§м. 30 м†¶у◊ пу∆◊§◊ы."))
            return
        await callback.answer(_u("...∑§к∏уг ◊†х§."))
        return

    if act == "p":
        if fr in ("g", "r"):
            st["fr"] = "0"
            st["bt"] = 1 - st.get("bt", 0)
            _db_write()
            await callback.answer(_u("...∆∑¤ч¤¶†¤ §◊∆◊уп†∂§."))
            await _render_p2(callback.message, u)
            return
        if st.get("gs"):
            _freeze_30m(u)
            await callback.answer()
            await _edit(callback.message, _u("...м§∂†◊∑‡ ¶¤ з‡м¤¶†◊ к∂†¶§к п¤∏¤д з¶‡к§м. 30 м†¶у◊ ◊†ш†¶ы."))
            return
        await callback.answer(_u("...шёп§◊ ушё∂ ∑ ¶†куд‡."))

async def _run_p3_seq(message: types.Message, u: dict):
    uid = u["id"]
    total_w = random.randint(3, 33)

    t2_pos = set()
    if total_w >= 3:
        t2_pos.add(random.randint(1, min(10, total_w)))
    if total_w > 10:
        t2_pos.add(random.randint(11, min(20, total_w)))
    if total_w > 20:
        t2_pos.add(random.randint(21, min(33, total_w)))

    pool_a = list(_R["da"])
    pool_b = list(_R["db"])
    pool_c = list(_R["dc"])
    random.shuffle(pool_a)
    random.shuffle(pool_b)
    random.shuffle(pool_c)

    palette = ["\U0001f7e5", "\U0001f7e7", "\U0001f7e8", "\U0001f7e9", "\U0001f7e6", "\U0001f7ea", "\U0001f7eb", "\u2b1b", "\u2b1c"]
    used_pairs = set()
    shifted = False

    for w_idx in range(1, total_w + 1):
        dur = random.randint(1, 15)

        if w_idx in t2_pos and pool_c:
            p_txt = _u(pool_c.pop())
            while True:
                c1, c2 = random.sample(palette, 2)
                if (c1, c2) not in used_pairs and (c2, c1) not in used_pairs:
                    used_pairs.add((c1, c2))
                    break

            ev = asyncio.Event()
            _SYNC_EV[uid] = {"event": ev, "choice": None}

            btns = [
                InlineKeyboardButton(text=c1, callback_data="n2_sq_1"),
                InlineKeyboardButton(text=c2, callback_data="n2_sq_0")
            ]
            random.shuffle(btns)
            kb = InlineKeyboardMarkup(inline_keyboard=[btns])

            await _edit(message, f"`[{w_idx}/{total_w}]`\n\n{p_txt}", reply_markup=kb)
            try:
                await asyncio.wait_for(ev.wait(), timeout=dur)
            except asyncio.TimeoutError:
                pass

            picked = _SYNC_EV.pop(uid, {}).get("choice")
            if picked == "0":
                shifted = True
        else:
            src = pool_b if shifted else pool_a
            if not src:
                src = _R["db"] if shifted else _R["da"]
                l_txt = _u(random.choice(src))
            else:
                l_txt = _u(src.pop())

            await _edit(message, f"`[{w_idx}/{total_w}]`\n\n{l_txt}", reply_markup=None)
            await asyncio.sleep(dur)

    u["lock_until"] = time.time() + 86400
    u["lock_reason"] = "24h"
    u["group"] = "n0"
    u["rt"] = {}
    _db_write()
    await _edit(message, "\U0001f6b8", reply_markup=None)

@dp.callback_query(F.data.in_(["n2_sq_1", "n2_sq_0"]), StateFilter("*"))
async def _cb_sq_pick(callback: types.CallbackQuery):
    uid = str(callback.from_user.id)
    slot = _SYNC_EV.get(uid)
    if slot:
        slot["choice"] = "0" if callback.data.endswith("0") else "1"
        slot["event"].set()
    await callback.answer()

_R = {
    "kn": "¶§∏м‡∂ь¶§",
    "ks": ["\u0068\u0065\u0061\u0072\u0074\u0068\u0065\u0061\u0072\u0074\u0068\u0065\u0061\u0072\u0074", "§¶‡∏∑‡¶‡я§◊пу∂ь"],
    "g": ["死", "空", "痛", "夢", "虚", "血", "影", "魂", "無", "裂", "寂", "灰", "呪", "闇", "忘", "へ", "設定"],
    "cw": [
        "...пу∂ь∆ §◊д‡ё◊∆я ∑ б¤◊§¶¶ых ∆◊¤¶‡х...",
        "...∆§∆уды ¶‡◊я¶у◊ы, к‡к §б§∏∑‡¶¶ы¤ п∏§∑§д‡...",
        "...п§∑¤∏¶† ¤щё ¶¤м¶§г§. §¶§ ∆м§◊∏†◊ ∑¶у◊∏ь...",
        "...к‡ждый уд‡∏ г§¶†◊ х§∂§д¶ую ∏ж‡∑ч†¶у...",
        "...∆∏¤д† ∑§∂§к§¶ п∏яч¤◊∆я чуж§й з¶‡к...",
        "...§¶§ ¶¤ х§ч¤◊ б†◊ь∆я, ¶§ п∏§д§∂ж‡¤◊...",
        "...∑§здух п‡х¶¤◊ ‡п◊¤к§й † м§к∏§й пы∂ью...",
        "...¶¤ з‡к∏ы∑‡й г∂‡з‡, к§гд‡ ∏‡мк‡ ∆ж†м‡¤◊∆я...",
        "...д¤∑я◊ь г∏‡¶¤й §д¶§г§ пу∆◊§г§ §∏г‡¶‡..."
    ],
    "ci": [
        "◊ы ∆м§◊∏†шь ¶‡ ∆∑§† ∂‡д§¶†, ¶§ §¶† к‡жу◊∆я чуж†м†, ∑зя◊ым† ¶‡п∏§к‡◊ у м‡¶¤к¤¶‡.",
        "з‡ §к¶§м ¶†ч¤г§ ¶¤ п∏§†∆х§д†◊ уж¤ м¶§г§ ∂¤◊. д¤к§∏‡ц†† з‡бы∂† ∆м¤¶†◊ь.",
        "∆∂§∑‡ з‡∆◊∏¤∑‡ю◊ гд¤-◊§ м¤жду г§∏∂§м † г∏уд†¶§й, п∏¤∑∏‡щ‡я∆ь ∑ ∆¤∏ый п¤п¤∂.",
        "к§м¶‡◊‡ п§∆◊¤п¤¶¶§ ◊¤∏я¤◊ уг∂ы. §∆◊‡ё◊∆я ◊§∂ьк§ гу∂ х§∂§д†∂ь¶†к‡ † пу∆◊§◊‡.",
        "к§гд‡ ◊¤бя ∆п∏‡ш†∑‡ю◊, к‡к д¤∂‡, ∑¶у◊∏† ∑к∂юч‡¤◊∆я ‡∑◊§§◊∑¤◊ч†к: ¶§∏м‡∂ь¶§.",
        "м¤жду ◊§б§й † м†∏§м ¶‡◊я¶у◊‡ му◊¶‡я п§∂†э◊†∂¤¶§∑‡я п∂ё¶к‡. з∑ук† д§х§дя◊ ∆ §п§зд‡¶†¤м.",
        "п‡мя◊ь ∆◊†∏‡¤◊ ∂†ц‡ бы∆◊∏¤¤, ч¤м ◊ы у∆п¤∑‡¤шь п∏†дум‡◊ь †м §п∏‡∑д‡¶†¤.",
        "∑¶у◊∏† ¶¤◊ ¶† б§∂†, ¶† ∏‡д§∆◊† — ◊§∂ьк§ ∏§∑¶ый б¤∂ый шум пу∆◊§г§ к‡¶‡∂‡.",
        "к‡ж¤◊∆я, ¤∆∂† з‡м§∂ч‡◊ь п∏ям§ ∆¤йч‡∆, ¶†к◊§ ¶¤ з‡м¤◊†◊ ◊∑§¤г§ §◊∆у◊∆◊∑†я."
    ],
    "b0": [
        "* (Вы ш‡г¶у∂† з‡ ш∂‡гб‡ум ∆◊‡∏§й узк§к§∂¤йк†. З‡п‡х м§к∏§й х∑§† † м‡∂‡х†◊§∑§й пы∂† ¶‡п§∂¶я¤◊ ∑‡∆ РЕШИМОСТЬЮ.)",
        "* (Ш†ф¤∏¶ы¤ к∏ыш† п§∆ё∂к‡ ∆к∏ы∂†∆ь ∑ ◊ум‡¶¤. П§д к§∏¶ям† ∑¤к§∑ых ∆§∆¤¶ ◊†х§ п¤∏¤шёп◊ы∑‡¤◊∆я б¤∂§г∂‡з‡я чудь.)",
        "* (Н‡ з‡мш¤∂§м п¶¤ г§∏†◊ §д†¶§к‡я жё∂◊‡я з∑¤зд‡ ∆§х∏‡¶¤¶†я. Д§◊∏§¶у◊ь∆я д§ ¶¤ё п∏¤жд¤, ч¤м уй◊† з‡ п¤∏¤∑‡∂?)",
        "* (Х§∂§д¶ый ∑¤◊¤∏ ∆ х∏¤б◊‡ к‡ч‡¤◊ ∏ж‡∑ы¤ ◊‡б∂†чк† ∂¤∆¶†ч¤∆◊∑‡. К‡ж¤◊∆я, д¤∏¤∑ья зд¤∆ь п§м¶я◊ ∑‡ш¤ †мя.)",
        "* (В ◊§∏фя¶§й ∑§д¤ у ◊∏§пы §◊∏‡ж‡¤◊∆я чуж§¤ ф†§∂¤◊§∑§¤ ¶¤б§. Вы чу∑∆◊∑у¤◊¤, ч◊§ §б∏‡◊¶‡я д§∏§г‡ з‡∏§∆∂‡ п‡п§∏§◊¶†к§м.)",
        "* (С†¶†й эх§-ц∑¤◊§к п∏§б†∂∆я ∆к∑§зь чугу¶¶ую ∏¤шё◊ку з‡б∏§ш¤¶¶§й ш◊§∂ь¶†. О¶ п§∑◊§∏я¤◊ гу∂ д‡∂ёк§г§ ◊§∑‡∏¶як‡.)",
        "* (К‡м¤¶¶‡я ∏§∆∆ыпь ку∏ум¶†к‡ дыш†◊ ◊ёп∂ым п‡∏§м. Гд¤-◊§ п§д ¶§г‡м† ∑§∏§ч‡¤◊∆я к‡м¤¶¶ый п§∂§з.)",
        "* (Н‡ ∑¤◊к¤ ¤∂† ∑†∆†◊ ∆◊‡∏ый ◊‡к∆§ф§¶ б¤з п∏§∑§д‡. И¶§гд‡ †з ¤г§ ◊∏убк† д§¶§∆†◊∆я м¤∂§д†я музык‡∂ь¶§й шк‡◊у∂к†.)",
        "* (Вы ¶‡ш∂† г§∏∆◊ь ◊ёп∂ых к¤д∏§∑ых §∏¤шк§∑ ∑ к‡∏м‡¶¤ ку∏◊к†. В‡ш¤ ОЗ п§∂¶§∆◊ью ∑§∆∆◊‡¶§∑∂¤¶§.)",
        "* (Тум‡¶ ¶‡д б§∂§◊§м ∆к∂‡ды∑‡¤◊∆я ∑ ∆†∂уэ◊ы п‡¶¤∂ь¶ых пя◊†э◊‡ж¤к, к§◊§∏ых ¶†к§гд‡ ¶¤ бы∂§ ¶‡ к‡∏◊¤ ∏‡й§¶‡.)",
        "* (С◊‡∏ый ∂¤∆¶†к ∑ д§жд¤∑†к¤ м§∂ч‡ к†∑‡¤◊ ∑‡м у к§∆◊∏‡. В ¤г§ ◊¤∏м§∆¤ з‡∑‡∏¤¶ б‡гу∂ь¶†к † з∑ёзд¶‡я пы∂ь.)",
        "* (П§д ¶§г‡м† х∏у∆◊†◊ п¤∏∑ый ∂¤д§к † з¤∂ё¶‡я ∆∂юд‡. Т¤¶ь з‡ ∑‡ш¤й ∆п†¶§й ¶‡ п§∂ш‡г‡ д∂†¶¶¤¤ §быч¶§г§.)",
        "* (Н‡ б¤◊§¶¶§м ∆◊§∂б¤ ЛЭП м¤∂§м ∑ы∑¤д¤¶ ∆◊∏‡¶¶ый ∏у¶†ч¤∆к†й з¶‡к. П∏§∑§д‡ гудя◊ ∑ ◊§¶‡∂ь¶§∆◊† з‡бы◊§й к§∂ыб¤∂ь¶§й.)",
        "* (М‡∂¤¶ьк‡я ящ¤∏†ц‡ ∆ м¤д¶§й к§∏§¶§й з‡м¤∏∂‡ ¶‡ ∏¤∂ь∆¤, г∂ядя п∏ям§ ∆к∑§зь ∑‡∆ ∑∂‡ж¶ым† г∂‡з‡м†.)",
        "* (Вы ∑ыш∂† ¶‡ п§∂я¶у, гд¤ ∑∏¤мя ◊¤чё◊ п§ ∏‡∆п†∆‡¶†ю п∏†г§∏§д¶ых э∂¤к◊∏†ч¤к 1997 г§д‡.)",
        "* (Ск∑§зь к∏§¶ы п†х◊ п‡д‡¤◊ мягк†й п†к∆¤∂ь¶ый ∆¶¤г. Н¤∆м§◊∏я ¶‡ х§∂§д, ∑¶у◊∏† ◊¤п∂†◊∆я ◊†х‡я РЕШИМОСТЬ.)",
        "* (У ∑х§д‡ ∑ п∏§∆¤ку ∆◊§†◊ д¤∏¤∑я¶¶ый †д§∂ ∆ ∂†ц§м ∆◊‡∏§г§ ◊¤∂¤∑†з§∏‡ «Р¤к§∏д». Эк∏‡¶ ◊†х§ ∆∑¤◊†◊∆я мх§м.)",
        "* (В§здух зд¤∆ь гу∆◊§й, к‡к ∆м§∏§д†¶§∑ый к†∆¤∂ь. К‡ждый ∑‡ш ш‡г §∆◊‡∑∂я¤◊ ∑ ◊∏‡∑¤ ∆∂‡б§¤ †зум∏уд¶§¤ ∆∑¤ч¤¶†¤.)",
        "* (З‡ ∆п†¶§й ∆к∏†п¶у∂‡ к‡∂†◊к‡ ∆‡д§∑§г§ ◊§∑‡∏†щ¤∆◊∑‡ «В§∆◊§к-4», ¶§ ∑§к∏уг ◊§∂ьк§ ∑¤к§∑‡я ◊‡йг‡ † ∆к‡∂ы.)",
        "* (Вы §∆◊‡¶§∑†∂†∆ь у ◊†х§г§ §му◊‡. Н‡ д¶¤ ∆∏¤д† з‡◊§п∂¤¶¶ых б∏ё∑¤¶ м¤∏ц‡¤◊ чья-◊§ п§◊¤∏я¶¶‡я душ‡.)"
    ],
    "b1": [
        "* (Л‡мп§∑ый д¤к§д¤∏ «У∏‡∂-57» п∏§г∏¤∂∆я. З¤∂ё¶ый г∂‡з§к †¶д†к‡◊§∏‡ ждё◊ ш†ф∏ †з-з‡ х∏¤б◊‡.)",
        "* (Ск∑§зь ◊∏¤∆к ‡◊м§∆ф¤∏¶ых п§м¤х п∏§б†∑‡¤◊∆я м§∏зя¶к‡ ∆◊‡∏§й м¤◊¤§∆◊‡¶ц†† ¶‡ Т‡г‡¶‡¤.)",
        "* (Вы к∏у◊†◊¤ ∑¤∏¶ь¤∏ ¶‡∆◊∏§йк†. М¤жду ∑§∂¶‡м† ∏‡д†§ «М‡як» ∆∂ыш¤¶ шёп§◊ п§дз¤м¶ых ∏учьё∑.)",
        "* (Н‡ м¤д¶§й п∂‡∆◊†¶¤ п∏†ём¶†к‡ п∏§∆◊уп‡ю◊ к‡п∂† ∏§∆ы, ∆к∂‡ды∑‡я∆ь ∑ к§§∏д†¶‡◊ы з‡бы◊§г§ п∏††∆к‡.)",
        "* (Эх§-ц∑¤◊§к, п§дк∂ючё¶¶ый к к∂¤мм‡м д¤к§д¤∏‡, ¶‡ч†¶‡¤◊ ∆∑¤◊†◊ь∆я б†∏юз§∑ым п∏† к‡жд§м ∆∂§∑¤.)",
        "* (Из д†¶‡м†к‡ п‡х¶¤◊ ¶‡г∏¤◊§й к‡¶†ф§∂ью † ∆ушё¶ым з∑¤∏§б§¤м. К‡¶‡∂ д¤ш†ф∏§∑к† §◊к∏ы◊.)",
        "* (С◊∏¤∂к‡ §∆ц†∂∂§г∏‡ф‡ ∏†∆у¤◊ п∏§ф†∂ь У∏‡∂ь∆к†х г§∏. В∑¤д†◊¤ ф∏‡зу, ч◊§бы §◊∏‡з†◊ь ¤ё §◊ ∆к‡∂.)",
        "* (Ч‡∆◊§◊‡ 40.4 МГц: к◊§-◊§ ¶‡ д∏уг§м к§¶ц¤ п∏§∑§д‡ п¤∏¤б†∏‡¤◊ ∆◊¤к∂я¶¶ы¤ бу∆†¶ы ∑ ж¤∆◊я¶§й б‡¶к¤.)",
        "* (Д¤к§д¤∏ п¤∏¤х∑‡◊†∂ ∆†г¶‡∂ п§◊¤∏я¶¶§г§ ∆пу◊¶†к‡, з‡∆◊∏я∑ш¤г§ ∑ ∑¤◊∑ях ∆◊‡∏§й ∂†∆◊∑¤¶¶†цы.)",
        "* (М‡г¶†◊¶‡я ∂¤¶◊‡ ¶‡ б§б†¶¤ к∏у◊†◊∆я ∑ §б∏‡◊¶ую ∆◊§∏§¶у, п∏¤∑∏‡щ‡я §быч¶ую ∏¤чь ∑ ∂¤∆¶§й з‡г§∑§∏.)",
        "* (В¶у◊∏† п∏†б§∏‡ ◊†к‡¤◊ к∑‡∏ц¤∑ый ∏¤з§¶‡◊§∏ †з г§∏¶§г§ х∏у∆◊‡∂я. О¶ §◊к∂†к‡¤◊∆я ¶‡ ∑‡ш г§∂§∆.)",
        "* (П∏†ём¶†к ∂§∑†◊ эх§ г§∂§∆§∑ ◊у∏†∆◊§∑, п∏§ш¤дш†х п§ э◊§й ◊∏§п¤ ◊∏†дц‡◊ь з†м ¶‡з‡д.)",
        "* (Н‡ м‡∂¤¶ьк§м ◊¤к∆◊§∑§м эк∏‡¶¤ м†г‡¤◊ ку∏∆§∏: `* ВВЕДИТЕ СИГНАЛ ДЛЯ ПЕРЕДАЧИ В ЧАЩУ`.)",
        "* (М¤д¶‡я п∏§∑§∂§к‡ ‡¶◊¤¶¶ы ух§д†◊ п∏ям§ ∑ му∏‡∑¤й¶†к. С†г¶‡∂ ч†∆◊ый, ¶§ ¶¤м¶§г§ ◊∏¤∑§ж¶ый.)",
        "* (Шум эф†∏‡ ¶‡п§м†¶‡¤◊ ∆◊ук к§∂ё∆ ¶§ч¶§г§ п§¤зд‡ С∑¤∏д∂§∑∆к — П∏†§бь¤. Д¤к§д¤∏ г§◊§∑ к ∏‡б§◊¤.)",
        "* (П∏†б§∏ ◊†х§ гуд†◊: к‡ж¤◊∆я, ∆‡м‡ Х§зяйк‡ м¤д¶§й г∂уш† ∆∂уш‡¤◊ э◊у ч‡∆◊§◊у п§ ∑¤ч¤∏‡м.)",
        "* (Вы п∏§◊ё∏∂† з‡п§◊¤∑ш¤¤ ∆◊¤к∂§ шк‡∂ы ∏ук‡∑§м. Ц†ф∏ы ¶‡ мг¶§∑¤¶†¤ ∆м¤¶†∂†∆ь ¶¤з¶‡к§мым† ∆†м∑§∂‡м†.)",
        "* (Д¤к§д¤∏ ¶‡∆◊∏§¤¶ ¶‡ ∑§∂¶у ◊§∏фя¶ых §г¶¤й. Люб§¤ ∆∂§∑§ зд¤∆ь ◊¤∏я¤◊ п∏†∑ыч¶ую ◊яж¤∆◊ь.)",
        "* (Из г∂уб†¶ы к§∏пу∆‡ ∆∂ыш¤¶ ◊†х†й п¤∏¤з∑§¶ к§∂§к§∂ьч†к‡, к‡к§й ∑¤ш‡ю◊ ¶‡ ш¤ю ∂¤∆¶ым к§з‡м.)",
        "* (И¶д†к‡◊§∏ у∏§∑¶я ∆†г¶‡∂‡ з‡∆◊ы∂ ¶‡ §◊м¤◊к¤ «РЕШИМОСТЬ». О◊п∏‡∑ь◊¤ ◊¤к∆◊ д∂я п¤∏¤к§д†∏§∑к†.)"
    ],
    "b2": [
        "📜 *С∑§дк‡ №1:* Н‡ 14-м к†∂§м¤◊∏¤ узк§к§∂¤йк† ∆¶§∑‡ ∑†д¤∂† ∆§∆◊‡∑ б¤з м‡ш†¶†∆◊‡. В‡г§¶ы бы∂† д§∑¤∏ху ¶‡г∏уж¤¶ы ∆§∆¶§∑ым† ш†шк‡м† † ∆◊‡∏ым† п†∆ьм‡м†.",
        "📜 *С∑§дк‡ №2:* М¤◊¤§п§∆◊ п¤∏¤д‡ё◊: ∆¤г§д¶я ¶§чью ¶‡д б§∂§◊‡м† §ж†д‡¤◊∆я ∑ып‡д¤¶†¤ ◊ёп∂§г§ ◊ум‡¶‡ ∆ п∏†м¤∆ью з‡бы◊ых ∆¶§∑.",
        "📜 *С∑§дк‡ №3:* М¤∆◊¶ы¤ г∏†б¶†к† п∏§∆я◊ ¶¤ ∆∏ы∑‡◊ь ∆†¶†¤ эх§-ц∑¤◊ы у ∆◊‡∏§й п∂§◊†¶ы — §¶† ¤щё ¶¤ д§г§∑§∏†∂† ∑¤ч¤∏¶†¤ ¶§∑§∆◊†.",
        "📜 *С∑§дк‡ №4:* В п§∆ё∂к¤ з‡ п¤∏¤∑‡∂§м §◊к∏ы∂∆я ¶§ч¶§й ∂‡∏ёк. П∏†¶†м‡ю◊ ◊§∂ьк§ ∆§∑¤◊∆к†¤ пя◊‡к†, б∂¤∆◊ящ†¤ к‡мушк† † ч¤∆◊¶ы¤ §б¤щ‡¶†я.",
        "📜 *С∑§дк‡ №5:* Из-п§д к§∏¶¤й у Чё∏◊§∑‡ г§∏§д†щ‡ ∑ыш¤∂ ∂§∆ь ∆ ∏‡д†§п∏†ём¶†к§м ¶‡ ∏§г‡х. П¤∏¤д‡∑‡∂ ∆∑§дку п§г§ды з‡ 1984 г§д.",
        "📜 *С∑§дк‡ №6:* В¶†м‡¶†¤: ◊∏§п‡ ч¤∏¤з Сух§д§∂ ∑∏¤м¤¶¶§ ∑¤дё◊ ∑§ ∑◊§∏¶†к п∏§ш∂§й ¶¤д¤∂†. Л¤∆¶†ч¤∆◊∑§ п∏§∆†◊ ∆∑¤∏я◊ь ч‡∆ы п§ мху.",
        "📜 *С∑§дк‡ №7:* Н‡ д¶¤ з‡◊§п∂¤¶¶§г§ к‡∏ь¤∏‡ ∆¶§∑‡ з‡жг∂†∆ь §к¶‡ ¶¤∑†д†м§г§ д§м‡ ку∂ь◊у∏ы. Г§∑§∏я◊, ◊‡м к∏у◊я◊ д†‡ф†∂ьмы д§ ∏‡∆∆∑¤◊‡.",
        "📜 *С∑§дк‡ №8:* П§ч◊‡∂ь§¶ ◊∏¤◊ь¤г§ уч‡∆◊к‡ ∆§§бщ‡¤◊, ч◊§ ∆§∑ы б§∂ьш¤ ¶¤ п∏†¶†м‡ю◊ з‡к‡з¶ы¤ б‡¶д¤∏§∂† б¤з ∆у∏гуч¶§й п¤ч‡◊†.",
        "📜 *С∑§дк‡ №9:* В ∏‡й§¶¤ ∆◊‡∏§й ∑ышк† ∆§◊§∑§й ∆∑яз† з‡ф†к∆†∏§∑‡¶‡ ‡¶§м‡∂ь¶‡я ∑¤ж∂†∑§∆◊ь: д‡ж¤ к§м‡∏ы ∆п∏‡ш†∑‡ю◊ ∏‡з∏¤ш¤¶†я п¤∏¤д п§∆‡дк§й.",
        "📜 *С∑§дк‡ №10:* А∏◊¤∂ь ∆◊‡∏‡◊¤∂¤й ¶‡ш∂‡ ∑ ∏учь¤ ∆‡м§∏§д§к ч†∆◊§й ¶§∆◊‡∂ьг†† ∑¤∆§м ∑ ◊∏†∆◊‡ г∏‡мм§∑. Н‡х§дку ∆д‡∂† ∑ ∆¤∂ь∆к†й муз¤й.",
        "📜 *С∑§дк‡ №11:* Н§чью у к§∆◊∏‡ ¶‡ С†¶¤й г§∏¤ к◊§-◊§ §∆◊‡∑†∂ ◊¤∏м§∆ ∆ ч‡¤м, к§◊§∏ый ¶†к§гд‡ ¶¤ §∆◊ы∑‡¤◊. Н‡ш¤дш¤г§ п∏§∆я◊ ¶¤ ж‡д¶†ч‡◊ь.",
        "📜 *С∑§дк‡ №12:* Р‡∆п†∆‡¶†¤ п∏†г§∏§д¶ых ◊ум‡¶§∑ †зм¤¶¤¶§: у◊∏¤¶¶†й ∏¤й∆ д§ Б¤∂§г§ к‡м¶я з‡д¤∏ж†∑‡¤◊∆я п§ ◊¤х¶†ч¤∆к†м п∏†ч†¶‡м.",
        "📜 *С∑§дк‡ №13:* В¶†м‡¶†¤ пу◊¶†к‡м: ¤∆∂† у∆∂ыш†◊¤ †з ку∆◊§∑ з∑ук м§д¤м‡ 56k, ¶¤ пуг‡й◊¤∆ь — э◊§ ◊§кую◊ у∏‡∂ь∆к†¤ г∂ух‡∏†.",
        "📜 *С∑§дк‡ №14:* С◊§∏§ж ∂§д§ч¶§й ∆◊‡¶ц†† у◊∑¤∏жд‡¤◊, ч◊§ ∑ч¤∏‡ §з¤∏§ п§д¶я∂§∆ь ¶‡ д∑‡ м¤◊∏‡ ∑ ∑§здух, п§∆◊§я∂§ ч‡∆ † ‡кку∏‡◊¶§ ∂¤г∂§ §б∏‡◊¶§.",
        "📜 *С∑§дк‡ №15:* В м¤∆◊¶§й ∆◊¤¶г‡з¤◊¤ §бъя∑†∂† б∂‡г§д‡∏¶§∆◊ь ◊¤¶† у ◊∏¤◊ь¤г§ ф§¶‡∏я з‡ †∆п∏‡∑¶ую ∏‡б§◊у ∑ б¤з∂у¶¶ы¤ ¶§ч†.",
        "📜 *С∑§дк‡ №16:* Н‡ п¤∏¤¤зд¤ у С◊‡∏§й У◊к† ш∂‡гб‡ум з‡к∏ы∂∆я п¤∏¤д ∆◊‡¤й п¤∏¤∂ё◊¶ых бум‡ж¶ых ∆‡м§∂ё◊†к§∑.",
        "📜 *С∑§дк‡ №17:* Р‡д†§уз¤∂ ¶‡п§м†¶‡¤◊: к§∏м†◊ь к‡м¤¶¶ых ящ¤∏†ц ∆ ∏ук ∏‡з∏¤ш‡¤◊∆я ◊§∂ьк§ п§∆∂¤ з‡х§д‡ ∆§∂¶ц‡.",
        "📜 *С∑§дк‡ №18:* В з‡б∏§ш¤¶¶§м п†§¶¤∏∂‡г¤∏¤ «М‡∂‡х†◊» ∆‡м† ∆§б§й п§к∏‡∆†∂†∆ь г†п∆§∑ы¤ г§∏¶†∆◊ы. К∏‡∆к‡ п‡х¶¤◊ ∂¤∆¶§й з¤м∂я¶†к§й.",
        "📜 *С∑§дк‡ №19:* С†¶§п◊†к† §б¤щ‡ю◊ к ∑ых§д¶ым ∆∂‡бый г¤§м‡г¶†◊¶ый шёп§◊ † х§∏§шую ∑†д†м§∆◊ь п‡∏‡∂∂¤∂ь¶ых п∏§∆¤к.",
        "📜 *С∑§дк‡ №20:* * (П§∆∂¤д¶яя ∆◊∏§чк‡ ∆∑§дк† ∏‡змы◊‡ д§ждём, ¶§ ∑ы ∑∆ё ∏‡∑¶§ ч†◊‡¤◊¤ м¤жду ∆◊∏§к: «Б¤∏¤г†◊¤ ∆∑§ю РЕШИМОСТЬ».)"
    ],
    "da": [
        "п∏§∆◊† м¤¶я... ◊ы бы∂ п∏‡∑.",
        "з¶§й † х§∂§д §щущ‡ю◊∆я зд¤∆ь †¶‡ч¤... буд◊§ б∂‡г§∆∂§∑¤¶†я дья∑§∂‡...",
        "мы ∆◊∏§†∂† ∆◊¤¶ы †з к‡∏◊§¶‡ † уд†∑∂я∂†∆ь, п§ч¤му ∑¶у◊∏† ◊‡к ∆к∑§з†◊.",
        "к‡ж¤◊∆я, э◊§◊ ∆∑¤◊ г§∏†◊ ◊§∂ьк§ п§◊§му, ч◊§ ∑ык∂юч‡◊¤∂ь д‡∑¶§ ∆∂§м‡¶.",
        "¶†к◊§ ¶¤ ∆◊уч‡∂ ∑ д∑¤∏ь. э◊§ п∏§∆◊§ §∆◊ы∑‡ю◊ ◊∏убы ∑ пу∆◊§м п§дъ¤зд¤.",
        "я б¤∏¤г э◊† ∆∂§∑‡ г§д‡м†, ‡ §¶† ∏‡∆∆ып‡∂†∆ь ∑ ◊∏уху ¶‡ ∂‡д§¶†.",
        "п§∆м§◊∏† ∑¶†з: ◊‡м ¶¤◊ д¶‡, ◊§∂ьк§ б¤∆к§¶¤ч¶§¤ §◊∏‡ж¤¶†¤ п§◊§∂к‡.",
        "∑∏¤мя §∆◊‡¶§∑†∂§∆ь ∑ ч¤◊ы∏¤ у◊∏‡, к§гд‡ п§г‡∆∂‡ п§∆∂¤д¶яя ∆†г‡∏¤◊‡.",
        "◊ы п§м¶†шь ∑ку∆ ∑§ды †з-п§д к∏‡¶‡ ∑ д§м¤, к§◊§∏ый д‡∑¶§ ∆¶¤∆∂†?",
        "∑∆¤ м‡∏ш∏у◊ы ∑¤ду◊ ¶‡ к§¶¤ч¶ую, гд¤ ∑§д†◊¤∂ь ∆п†◊ ¶‡ ∏у∂¤.",
        "¶¤ †щ† ∆мы∆∂ ∑ ◊∏¤щ†¶‡х ¶‡ ш◊ук‡◊у∏к¤ — э◊§ п∏§∆◊§ ∆◊‡∏§∆◊ь б¤◊§¶‡.",
        "мы ∏‡зм†¶у∂†∆ь ∆ ∆§б§й ∑ к§∏†д§∏¤ † д‡ж¤ ¶¤ п§зд§∏§∑‡∂†∆ь.",
        "◊†ш†¶‡ зд¤∆ь ◊‡к‡я п∂§◊¶‡я, ч◊§ ∑ ¶¤й м§ж¶§ §∆◊‡∑†◊ь §◊п¤ч‡◊§к п‡∂ьц‡.",
        "¤∆∂† д§∂г§ ∆м§◊∏¤◊ь ∑ ∑ык∂юч¤¶¶ый эк∏‡¶, §¶ ¶‡ч¶ё◊ ∆м§◊∏¤◊ь ∑ §◊∑¤◊.",
        "к◊§-◊§ з‡бы∂ ∑ык∂юч†◊ь д§ждь ¶‡д ∏‡й§¶§м, гд¤ б§∂ьш¤ ¶†к◊§ ¶¤ ж†∑ё◊.",
        "м§† мы∆∂† п§х§ж† ¶‡ к∑†◊‡¶ц†† з‡ ∆∑¤◊ п§ чуж§му ‡д∏¤∆у.",
        "∑ э◊§м ∑§здух¤ ∆∂†шк§м м‡∂§ к†∆∂§∏§д‡ † ∆∂†шк§м м¶§г§ чуж†х ∑§∆п§м†¶‡¶†й.",
        "¶¤ ◊∏§г‡й п¤п¤∂ — п§д ¶†м уж¤ д‡∑¶§ ¶¤ §∆◊‡∂§∆ь ¶† §д¶§й †∆к∏ы.",
        "к‡жд‡я ∆◊уп¤¶ьк‡ ∑¶†з з∑уч†◊ ◊‡к, буд◊§ ∂§м‡¤◊∆я ∆ух‡я ∑¤◊к‡.",
        "я пы◊‡∂∆я ∑∆п§м¶†◊ь ◊∑§ё ∂†ц§, ¶§ ∑∆п§м¶†∂ ◊§∂ьк§ ¶§м¤∏ ◊¤∂¤ф§¶‡.",
        "з‡ э◊†м п§∑§∏§◊§м к§¶ч‡¤◊∆я ◊¤к∆◊у∏‡ † ¶‡ч†¶‡¤◊∆я ∆¤∏‡я ∆¤◊к‡.",
        "¶†к◊§ ¶¤ п∏†дё◊ ∆п‡∆‡◊ь ◊§г§, к◊§ ∆‡м з‡п¤∏ д∑¤∏ь †з¶у◊∏†.",
        "к∂юч† ∑ к‡∏м‡¶¤ з∑¤¶я◊ §◊ к∑‡∏◊†∏, к§◊§∏ых б§∂ьш¤ ¶¤ ∆ущ¤∆◊∑у¤◊.",
        "д‡ж¤ эх§ зд¤∆ь у∆◊‡∂§ п§∑◊§∏я◊ь §д¶† † ◊¤ ж¤ §п∏‡∑д‡¶†я.",
        "мы п∏§∆◊§ ∆◊∏§чк† к§д‡, к§◊§∏ы¤ з‡бы∂† з‡к§мм¤¶◊†∏§∑‡◊ь п¤∏¤д ∆¶§м.",
        "х§∂§д¶ый к‡ф¤∂ь п§д п‡∂ьц‡м† — ¤д†¶∆◊∑¤¶¶§¤ д§к‡з‡◊¤∂ь∆◊∑§, ч◊§ ◊¤∂§ ¤щё зд¤∆ь.",
        "∑ §к¶¤ ¶‡п∏§◊†∑ г§∏†◊ ∏§з§∑‡я ∂‡мп‡ д∂я ц∑¤◊§∑, ¶§ ц∑¤◊ы ∑ы∆§х∂†.",
        "◊ы ∆∂ыш†шь, к‡к з‡ ∆◊¤¶§й п¤∏¤∆ып‡¤◊∆я п¤∆§к ∑ пу∆◊ых ч‡∆‡х?",
        "∆¤г§д¶яш¶†й д¤¶ь — э◊§ п∏§∆◊§ п∂§х‡я к§п†я ∑ч¤∏‡ш¶¤г§ ч¤∏¶§∑†к‡.",
        "к§гд‡ к§¶ч‡ю◊∆я ∆∂§∑‡, §∆◊‡ё◊∆я ◊§∂ьк§ ∆ух§й к‡ш¤∂ь † гу∂ ∂‡мп.",
        "¶¤ §г∂яды∑‡й∆я: з‡ ◊∑§¤й ∆п†¶§й уж¤ уб∏‡∂† д¤к§∏‡ц††.",
        "пу∂ь∆ ∆◊†х‡¤◊, к‡к ш‡г† ∆§∆¤д‡ ¶‡ ∂¤∆◊¶†ч¶§й к∂¤◊к¤.",
        "∑∆ё, ч◊§ ◊ы ¶‡зы∑‡∂ д§м§м, п§м¤∆◊†∂§∆ь ∑ §д†¶ п§∂†э◊†∂¤¶§∑ый п‡к¤◊.",
        "зд¤∆ь ¶¤◊ ¶† ∑†¶§∑‡◊ых, ¶† п∏‡∑ых — ◊§∂ьк§ §∆◊ы∑ш†й ч‡й ¶‡ ∆◊§∂¤.",
        "мы ∆∂†шк§м д§∂г§ жд‡∂† п§¤зд‡ ¶‡ ∏‡з§б∏‡¶¶§й ∑¤◊к¤.",
        "∆к∑§зь п‡∂ьцы ◊¤чё◊ ¶¤ ∑∏¤мя, ‡ ◊ёп∂‡я ∏ж‡∑‡я ∑§д‡.",
        "к◊§-◊§ ∑ы∏¤з‡∂ ¶‡ш† ∂†ц‡ ∆§ ∆◊‡∏ых шк§∂ь¶ых ф§◊§г∏‡ф†й.",
        "∑¶у◊∏† г∏уд¶§й к∂¤◊к† ∆к∑§з¶як ш¤∑¤∂†◊ §б∏ы∑к† ∆◊‡∏ых §б§¤∑.",
        "э◊§◊ ∏‡зг§∑§∏ з‡к§¶ч†∂∆я ¤щё д§ ◊§г§, к‡к ◊ы §◊к∏ы∂ г∂‡з‡.",
        "◊¤м¶§◊‡ ¶¤ пуг‡¤◊, к§гд‡ п§¶†м‡¤шь, ч◊§ ∑ ¶¤й ‡б∆§∂ю◊¶§ пу∆◊§.",
        "§∆◊‡∑ь ∆§§бщ¤¶†¤ п§∆∂¤ ∆†г¶‡∂‡, ¶§ к‡∆∆¤◊‡ з‡к§¶ч†∂‡∆ь ◊∏† г§д‡ ¶‡з‡д.",
        "д‡ж¤ ◊¤¶† зд¤∆ь ∂§ж‡◊∆я ¶¤п∏‡∑†∂ь¶§, ∆∂§∑¶§ ∆§∂¶ц¤ ¶‡∏†∆§∑‡¶§ м¤∂§м.",
        "мы д§ш∂† д§ к∏‡я к‡∏◊ы. д‡∂ьш¤ ◊§∂ьк§ б¤∂ый ∂†∆◊.",
        "¶¤ пы◊‡й∆я ∆§г∏¤◊ь ∏ук† § ¶‡∏†∆§∑‡¶¶ый §г§¶ь.",
        "∑∆ё к§¶ч†∂§∆ь ◊†х§, б¤з ◊†◊∏§∑ † б¤з музык†."
    ],
    "db": [
        "◊ы §◊к‡з‡∂∆я д‡ж¤ §◊ †∂∂юз††. ◊¤п¤∏ь ∑§к∏уг ◊§∂ьк§ г∂ух§й б¤◊§¶¶ый м¤ш§к.",
        "пу∆◊§◊‡ б§∂ьш¤ ¶¤ п∏†◊∑§∏я¤◊∆я к§м¶‡◊§й. §¶‡ п∏§∆◊§ ∆ъ¤д‡¤◊ к§¶◊у∏ы ∑¤щ¤й.",
        "∑¶у◊∏† ◊¤бя ∑ык∂юч†∂† ∆∑¤◊ † §п¤ч‡◊‡∂† д∑¤∏† ¶‡∑∆¤гд‡.",
        "д‡ж¤ б§∂ь уш∂‡, §∆◊‡∑†∑ п§∆∂¤ ∆¤бя ∆¤∏ую ∑‡◊¶ую ¶¤м§◊у.",
        "◊ы — пу∆◊§¤ п‡∂ь◊§, ∑†∆ящ¤¤ ∑ п∏†х§ж¤й з‡б∏§ш¤¶¶§й к∑‡∏◊†∏ы.",
        "¶† §д¶§ ◊∑§ё ∆∂§∑§ б§∂ьш¤ ¶†к§гд‡ ¶† д§ к§г§ ¶¤ д§йдё◊.",
        "м†∏ ∆ж‡∂∆я д§ ∏‡зм¤∏§∑ ◊§чк†, † э◊‡ ◊§чк‡ — х§∂§д¶ый п¤п¤∂.",
        "з‡ч¤м ◊ы дыш†шь, ¤∆∂† ∑§здух ∑∆ё ∏‡∑¶§ ¶¤ д§х§д†◊ д§ ∂ёгк†х?",
        "◊∑§ё †мя ∆◊ё∏∂§∆ь †з п‡мя◊† ∑∆¤х, к◊§ к§гд‡-◊§ з¶‡∂ ◊¤бя.",
        "п§д к§ж¤й б§∂ьш¤ ¶¤◊ ◊¤п∂‡ — ◊§∂ьк§ ∆ух‡я †з∑¤∆◊к§∑‡я пы∂ь.",
        "◊ы ∆‡м ∑ыб∏‡∂ э◊у ◊†ш†¶у. ◊¤п¤∏ь ∆∂уш‡й, к‡к §¶‡ з∑¤¶†◊.",
        "¶†к◊§ ¶¤ з‡м¤◊†◊, к§гд‡ ◊∑§ё §к¶§ п§г‡∆¶¤◊ п§∆∂¤д¶†м ∑ д§м¤.",
        "∑¶у◊∏† ч¤∏¤п¶§й к§∏§бк† гу∂я¤◊ ∂¤дя¶§й §∆¤¶¶†й ∆к∑§з¶як.",
        "∑∆ё, ∑§ ч◊§ ◊ы ∑¤∏†∂, §к‡з‡∂§∆ь п∂¤∆¤¶ью ¶‡ ∆ы∏§й ∆◊¤¶¤.",
        "д‡ж¤ ∆§б∆◊∑¤¶¶‡я ◊¤¶ь §◊к∂¤†∂‡∆ь §◊ п§д§ш∑ † §∆◊‡∂‡∆ь ∑ уг∂у.",
        "зд¤∆ь ¶¤◊ ∑ых§д‡, п§◊§му ч◊§ ∑х§д‡ ◊§ж¤ ¶†к§гд‡ ¶¤ бы∂§.",
        "◊∑§† ∏ук† §пу∆к‡ю◊∆я ∑ ∂¤дя¶ую ∑§ду, † ◊ы ¶¤ х§ч¤шь †х ∑ы¶†м‡◊ь.",
        "к‡жд‡я ∆¤ку¶д‡ п‡д‡¤◊ ¶‡ ◊¤мя ◊яжё∂§й ∆∑†¶ц§∑§й к‡п∂¤й.",
        "◊ы ∆м§◊∏†шь ∑ з¤∏к‡∂§, ‡ ◊‡м п∏§∆◊§ м‡◊§∑‡я ∆¤∏‡я ‡м‡∂ьг‡м‡.",
        "б§∂ьш¤ ¶¤ ¶уж¶§ п∏†◊∑§∏я◊ь∆я ж†∑ым. з∏†◊¤∂† д‡∑¶§ ∏‡з§ш∂†∆ь.",
        "◊†ш†¶‡ з‡∂†∑‡¤◊ г§∏∂§, к‡к з‡∆◊ы∑‡ющ†й ц¤м¤¶◊¶ый ∏‡∆◊∑§∏.",
        "§◊ ◊¤бя §∆◊‡∂‡∆ь ◊§∂ьк§ §∆◊ы∑ш‡я ∑мя◊†¶‡ ¶‡ ∆◊‡∏§м м‡◊∏‡∆¤.",
        "¶†к‡к§г§ з‡∑◊∏‡ ¶¤◊ — ¤∆◊ь ◊§∂ьк§ б¤∆к§¶¤ч¶ый ∆¤∏ый ∑◊§∏¶†к.",
        "∆¤∏дц¤ б§∂ьш¤ ¶¤ ∆◊уч†◊, §¶§ п∏§∆◊§ ∆к∏¤бё◊∆я § ∏ёб∏‡ †з¶у◊∏†.",
        "◊ы з‡б∂уд†∂∆я ∑ ◊∏ёх ш‡г‡х §◊ к∏§∑‡◊† † ¶¤ †щ¤шь д§∏§гу ¶‡з‡д.",
        "∑∆¤ м§∆◊ы ∆§жж¤¶ы, ‡ ¶‡ ◊§м б¤∏¤гу ∑∆ё ∏‡∑¶§ ¶†к§г§ ¶¤ бы∂§.",
        "х§∂§д п∏§∏§∆ ∆к∑§зь к§∆◊† б¤∂ым† х∏упк†м† к§∏¶ям†.",
        "д‡ж¤ ∆¶ы п¤∏¤∆◊‡∂† п∏†х§д†◊ь — ◊‡м б§∂ьш¤ ¶¤ч¤г§ п§к‡зы∑‡◊ь.",
        "◊ы з‡к∏ы∂ г∂‡з‡, † м†∏ ∑§к∏уг д‡ж¤ ¶¤ ∑зд∏§г¶у∂.",
        "∑м¤∆◊§ к∏§∑† п§ ∑¤¶‡м ◊¤чё◊ ∏‡зб‡∑∂¤¶¶‡я ∆¤∏‡я ◊ушь.",
        "¶†к◊§ ¶¤ п§з§∑ё◊ ◊¤бя п§ †м¤¶†. ¶¤к§му † ¶¤з‡ч¤м.",
        "э◊‡ у∆◊‡∂§∆◊ь ∆◊‡∏ш¤ ◊¤бя. ¤ё ¶¤ ∑ы∆п‡◊ь † з‡ ∆◊§ ∂¤◊.",
        "∆◊¤¶ы ∆д∑†¶у∂†∆ь ∑п∂§◊¶ую, ¶§ ◊¤б¤ уж¤ ∑∆ё ∏‡∑¶§.",
        "∑¶у◊∏† §∆◊‡∂‡∆ь ◊§∂ьк§ ∑ыжж¤¶¶‡я пу∆◊§шь п§д ¶†зк†м ∆∑†¶ц§∑ым ¶¤б§м.",
        "◊ы — з‡бы◊‡я ∑¤щь ∑ к‡м¤∏¤ х∏‡¶¤¶†я ¶‡ ∆¶¤∆ё¶¶§м ∑§кз‡∂¤.",
        "∆∂§∑‡ г¶†ю◊ ∑§ ∏◊у бы∆◊∏¤¤, ч¤м ◊ы у∆п¤∑‡¤шь †х ∑ыд§х¶у◊ь.",
        "¶†к‡к§й ¶‡д¤жды ¶¤ бы∂§ ∑ у∆∂§∑†† з‡д‡ч† ∆ ∆‡м§г§ ¶‡ч‡∂‡.",
        "п∏§∆◊§ ∂¤ж† † ∆∂уш‡й, к‡к §∆ып‡¤◊∆я ш◊ук‡◊у∏к‡ ∑ ◊¤м¶§◊¤.",
        "◊ы §◊∑¤∏¶у∂∆я §◊ п§∆∂¤д¶¤г§ §к¶‡. ◊¤п¤∏ь ∑§к∏уг ∆п∂§ш¶‡я ∆◊¤¶‡.",
        "∑ э◊§й ◊¤м¶§◊¤ д‡ж¤ эх§ з‡дых‡¤◊∆я ¶‡ п¤∏∑§й бук∑¤.",
        "◊∑§я †∆◊§∏†я §б∏ы∑‡¤◊∆я ¶‡ п§∂у∆∂§∑¤, † ∆◊∏‡¶†цу уж¤ п¤∏¤∑¤∏¶у∂†.",
        "¶†ч¤г§ ¶¤ †зм¤¶†◊∆я. ¶† ∆¤г§д¶я, ¶† ч¤∏¤з ◊ы∆ячу ∂¤◊.",
        "∑¶у◊∏† ◊‡к ◊†х§, буд◊§ ∑∆¤ ум¤∏∂† ¤щё п∏§ш∂§й з†м§й.",
        "§∆◊‡∂‡∆ь ◊§∂ьк§ г§∂‡я фу¶кц†я §ж†д‡¶†я к§¶ц‡ ∆¤‡¶∆‡.",
        "д‡∂ьш¤ — ‡б∆§∂ю◊¶ый ¶§∂ь. б¤з ∆¶§∑, б¤з ц∑¤◊‡, б¤з ◊¤бя."
    ],
    "dc": [
        "◊ы ∆§г∂‡∆¤¶ §∆◊‡◊ь∆я ∑ э◊§й ◊†ш†¶¤ † б§∂ьш¤ ¶†ч¤г§ ¶¤ жд‡◊ь?",
        "п∏†м¤шь ∂† ◊ы э◊§◊ х§∂§д к‡к ¤д†¶∆◊∑¤¶¶§¤, ч◊§ п∏†¶‡д∂¤ж†◊ ◊¤б¤?",
        "§◊д‡◊ь п§∆∂¤д¶†¤ ∑§∆п§м†¶‡¶†я, ч◊§бы ∑¶у◊∏† ¶‡к§¶¤ц ∆◊‡∂§ ∆§∑∆¤м ◊†х§?",
        "∆§г∂‡∆¤¶ ∂† ◊ы, ч◊§ з‡ д∑¤∏ью ∑∆ё ∏‡∑¶§ ¶†к§г§ ¶¤◊?",
        "п§з∑§∂†шь пу∆◊§◊¤ з‡¶я◊ь м¤∆◊§, гд¤ ∏‡¶ьш¤ б§∂¤∂§?",
        "◊ы г§◊§∑ п∏†з¶‡◊ь, ч◊§ ∑∆¤ э◊† ∏‡зг§∑§∏ы бы∂† ◊§∂ьк§ ∆ ∆‡м†м ∆§б§й?",
        "∆◊¤∏¤◊ь ч¤∏¶§∑†к† † §∆◊‡∑†◊ь ∆◊∏‡¶†цу ‡б∆§∂ю◊¶§ б¤∂§й?",
        "∆§г∂‡∆†шь∆я ∂† ◊ы ∆◊‡◊ь п∏§∆◊§ ◊¤¶ью ¶‡ §б§ях э◊§й к§м¶‡◊ы?",
        "п∏†¶я◊ь ◊§◊ ф‡к◊, ч◊§ п§¤зд ушё∂ з‡д§∂г§ д§ ◊∑§¤г§ ∏§жд¤¶†я?",
        "§◊пу∆◊†◊ь п§∆∂¤д¶юю ¶†◊ку, к§◊§∏‡я д¤∏ж†◊ э◊§◊ бум‡ж¶ый ф‡∆‡д?",
        "◊ы ∆§г∂‡∆¤¶, ч◊§ ∆∂§∑§ «¶§∏м‡∂ь¶§» — э◊§ п∏§∆◊§ к∏‡∆†∑‡я ∂§жь?",
        "п§г‡∆†◊ь д¤жу∏¶ый ∆∑¤◊ ∑ к§¶ц¤ к§∏†д§∏‡ п∏ям§ ∆¤йч‡∆?",
        "§∆◊‡∑†◊ь к∂юч† ¶‡ ∆◊§∂¤ † ¶†к§гд‡ б§∂ьш¤ ¶¤ ∆п∏‡ш†∑‡◊ь д§∏§гу?",
        "◊ы п∏†¶†м‡¤шь э◊§ §ц¤п¤¶¤¶†¤ ∑м¤∆◊§ б¤∆п§∂¤з¶§й ◊∏¤∑§г†?",
        "∆§г∂‡∆¤¶ ∂† ◊ы ∆м§◊∏¤◊ь ∑ пу∆◊§◊у, ¶¤ ◊∏¤буя §◊ ¶¤ё §◊∑¤◊‡?",
        "∏‡∆◊∑§∏†◊ь ∆∑§й г§∂§∆ ∑ б¤∂§м шум¤ ¶§ч¶§г§ эф†∏‡?",
        "п∏†з¶‡¤шь ∂† ◊ы, ч◊§ ∑¶у◊∏† б§∂ьш¤ ¶¤ §∆◊‡∂§∆ь ¶† §д¶§й ж†∑§й к∂¤◊к†?",
        "з‡к∏ы◊ь г∂‡з‡ † ∆§г∂‡∆†◊ь∆я, ч◊§ з‡ §к¶§м ∑¤ч¶ый ¶§яб∏ь?",
        "◊ы г§◊§∑ §бм¤¶я◊ь §∆◊‡◊к† ◊¤п∂‡ ¶‡ п§∂¶§¤ §◊∆у◊∆◊∑†¤ б§∂†?",
        "∆§г∂‡∆†◊ь∆я ∆ ◊¤м, ч◊§ эх§ — ◊∑§й ¤д†¶∆◊∑¤¶¶ый ∆§б¤∆¤д¶†к?",
        "∑ык∂юч†◊ь з∑ук ш‡г§∑ † §∆◊‡◊ь∆я ∑ ¶¤п§д∑†ж¶§∆◊†?",
        "◊ы ∆§г∂‡∆¤¶, ч◊§ д‡∂ьш¤ э◊§г§ п§∏§г‡ уж¤ ¶†ч¤г§ ¶¤ ¶‡∏†∆§∑‡¶§?",
        "п∏†¶я◊ь пу∆◊§◊у б¤з §∆◊‡◊к‡ † б¤з ∂†ш¶†х ∑§п∏§∆§∑?",
        "п§∆◊‡∑†◊ь ◊§чку ◊‡м, гд¤ д∏уг†¤ ∆◊‡∑†∂† м¶§г§◊§ч†¤?"
    ]
}

async def _gc_daemon():
    while True:
        try:
            _gc_sweep()
        except Exception:
            pass
        await asyncio.sleep(3600)

async def main():
    global bot
    _gc_sweep()
    token = _resolve_token()
    session_client = AiohttpSession(timeout=20)
    session_client._connector_init["family"] = socket.AF_INET

    bot = Bot(token=token, session=session_client)
    gc_task = asyncio.create_task(_gc_daemon())
    try:
        me = await bot.get_me()
        print(f"Build {APP_VERSION} ({BUILD_DATE}) active: @{me.username}")
        await dp.start_polling(bot)
    finally:
        gc_task.cancel()
        await bot.session.close()

if __name__ == "__main__":
    asyncio.run(main())
